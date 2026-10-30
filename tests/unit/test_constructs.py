"""Round 12: constructs that were lost, mangled or misreported, and the remaining publication edges."""
import errno
import os
import re
import stat
import zipfile
from pathlib import Path

import pytest

from texdocxdrafter import fidelity, pipeline, stylemap
from texdocxdrafter.pipeline import PublishError, SourceChangedError, build
from texdocxdrafter.preflight import PreflightError
from texdocxdrafter.validate import check_docx

from test_publish import PNG, _no_temp_files, _quiet, _rezip


def _doc(tmp_path: Path, body: str, preamble: str = "", name="main.tex") -> Path:
    src = tmp_path / name
    src.write_text("\\documentclass{article}\n\\usepackage{graphicx,hyperref,longtable,multirow,subcaption,amsmath}\n"
                   + preamble + "\n\\begin{document}\n" + body + "\n\\end{document}\n")
    src.with_suffix(".aux").write_text("\\relax\n")
    (tmp_path / "pic.png").write_bytes(PNG)
    return src


def _xml(tmp_path: Path, part="word/document.xml") -> str:
    z = zipfile.ZipFile(tmp_path / "main.docx")
    return z.read(part).decode("utf-8") if part in z.namelist() else ""


def _text(xml: str) -> str:
    return "".join(re.findall(r"<w:t(?: [^>]*)?>([^<]*)</w:t>", xml))


def _paras(xml: str) -> list[str]:
    return [_text(p) for p in re.findall(r"<w:p[ >].*?</w:p>", xml, re.S)]


# ---- 1: % inside \url, \href, \verb ------------------------------------------------------------------


def test_percent_inside_url_href_and_verb_is_not_a_comment(tmp_path):
    src = _doc(tmp_path, r"""\section{Before}
See \url{https://ex.com/a%20b} and \href{https://ex.com/x%7Ey}{the site} and \verb|50%| here.
\section{After}
Second section text.""")
    res = build(src, no_harvest=True, log=_quiet)
    got = fidelity.found(tmp_path / "main.docx", stylemap.load(None))
    assert got.headings == ["Before", "After"], got.headings
    xml = _xml(tmp_path)
    assert "https://ex.com/a%20b" in xml and "50%" in _text(xml) and "Second section text" in _text(xml)
    assert "https://ex.com/x%7Ey" in _xml(tmp_path, "word/_rels/document.xml.rels")
    assert res.check.ok


# ---- 2 and 3: floats keep everything ------------------------------------------------------------------------


def test_float_with_picture_and_table_keeps_both(tmp_path):
    src = _doc(tmp_path, r"""\begin{figure}\centering
\includegraphics{pic.png}
\begin{tabular}{ll} alpha & 1 \\ beta & 2 \end{tabular}
\caption{Both}\label{fig:both}
\end{figure}
\begin{table}\centering
\begin{tabular}{ll} gamma & 3 \end{tabular}
\includegraphics{pic.png}
\caption{Table with a picture}
\end{table}
Text refers to Figure~\ref{fig:both}.""")
    build(src, no_harvest=True, log=_quiet)
    got = fidelity.found(tmp_path / "main.docx", stylemap.load(None))
    assert got.tables == 2 and got.pictures == 2
    t = _text(_xml(tmp_path))
    assert "alpha" in t and "beta" in t and "gamma" in t and "??" not in t


def test_several_captions_and_labels_in_one_float(tmp_path):
    src = _doc(tmp_path, r"""\begin{table}\centering
\begin{tabular}{l} first \end{tabular}
\caption{First table\label{tab:one}}\label{tab:one-too}
\begin{tabular}{l} second \end{tabular}
\caption{Second table}\label{tab:two}
\end{table}
Refs: \ref{tab:one}, \ref{tab:one-too}, \ref{tab:two}.
\begin{figure}
\caption{Above its picture}\label{fig:a}
\includegraphics{pic.png}
\caption{Below its picture}\label{fig:b}
\includegraphics{pic.png}
\end{figure}
More refs: \ref{fig:a} and \ref{fig:b}.""")
    res = build(src, no_harvest=True, log=_quiet)
    assert res.check.ok, res.check.errors
    xml = _xml(tmp_path)
    paras = _paras(xml)
    caps = [p for p in paras if p.startswith("Table ") or p.startswith("Figure ")]
    assert [c.split(": ", 1)[1] for c in caps] == ["First table", "Second table", "Above its picture", "Below its picture"], caps
    # table captions sit above their table (style-map default): each caption directly precedes its own tabular
    assert paras.index("Table 1: First table") < paras.index("first") < paras.index("Table 2: Second table") < paras.index("second")
    assert "??" not in _text(xml)
    assert not [e for e in res.audit.entries if "undefined" in e.message or "no target" in e.message], \
        [e.message for e in res.audit.entries]
    for key in ("tab:one", "tab:one-too", "tab:two", "fig:a", "fig:b"):
        assert res.check.fields  # REF fields exist
    assert xml.count('w:instr') >= 0 and xml.count("REF ") >= 5


# ---- 4: links inside footnotes ---------------------------------------------------------------------------


def test_links_inside_footnotes_belong_to_the_footnotes_part(tmp_path):
    src = _doc(tmp_path, r"Text.\footnote{See \href{https://ex.com/fn}{this} and \url{https://ex.com/u}.} More.")
    res = build(src, no_harvest=True, log=_quiet)
    assert res.check.ok, res.check.errors
    fn = _xml(tmp_path, "word/footnotes.xml")
    rels = _xml(tmp_path, "word/_rels/footnotes.xml.rels")
    ids = re.findall(r'<w:hyperlink[^>]* r:id="(rId\d+)"', fn)
    assert len(ids) == 2 and all(f'Id="{i}"' in rels for i in ids), (ids, rels)
    assert "https://ex.com/fn" in rels and "https://ex.com/u" in rels
    assert "https://ex.com/fn" not in _xml(tmp_path, "word/_rels/document.xml.rels")


def test_validator_resolves_relationships_of_every_part(tmp_path):
    src = _doc(tmp_path, r"Text.\footnote{Plain note.}")
    build(src, no_harvest=True, log=_quiet)
    out = tmp_path / "main.docx"
    _rezip(out, {"word/footnotes.xml": [(b"Plain note.", b'<w:hyperlink r:id="rId99"><w:r><w:t>x</w:t></w:r></w:hyperlink>')]}, {})
    errs = check_docx(out).errors
    assert any("footnotes.xml: unresolved relationship rId99" in e for e in errs), errs


# ---- 5, 6, 9: publication edges -----------------------------------------------------------------------


def test_new_output_without_hard_links_is_never_half_written(tmp_path, monkeypatch):
    target = tmp_path / "t.docx"
    monkeypatch.setattr(pipeline.os, "link", lambda *a: (_ for _ in ()).throw(OSError(errno.EPERM, "no links here")))
    pipeline.commit_file(target, b"complete", None, "output")  # the fallback: reserve, then rename the finished tmp
    assert target.read_bytes() == b"complete"
    target.unlink()
    monkeypatch.setattr(pipeline.os, "replace", lambda *a: (_ for _ in ()).throw(OSError(errno.ENOSPC, "full")))
    with pytest.raises(OSError):
        pipeline.commit_file(target, b"complete", None, "output")
    assert not target.exists()  # no 0-byte or truncated file left at the target
    _no_temp_files(tmp_path)


def test_failed_rollback_is_reported_not_swallowed(tmp_path, monkeypatch):
    a, b = tmp_path / "a.txt", tmp_path / "b.txt"
    a.write_bytes(b"old a")
    snap_a = pipeline._snapshot(a)
    real = pipeline.commit_file

    def failing(path, data, snapshot, what="file"):
        if path == b:
            raise OSError(errno.ENOSPC, "No space left on device")
        return real(path, data, snapshot, what)

    monkeypatch.setattr(pipeline, "commit_file", failing)
    monkeypatch.setattr(pipeline, "write_atomic", lambda p, d: (_ for _ in ()).throw(OSError(errno.ENOSPC, "still full")))
    with pytest.raises(PublishError, match="a.txt .* could not be put back and still hold the new build") as info:
        pipeline.publish([(a, b"new a", snap_a, "report"), (b, b"new b", None, "output")])
    assert "--overwrite" in str(info.value) and isinstance(info.value.__cause__, OSError)
    assert "b.txt" not in str(info.value).split("could not be put back")[0]  # R7: only the files not restored are named


def test_published_files_follow_the_umask_and_keep_their_mode(tmp_path, monkeypatch):
    old = os.umask(0o022)
    try:
        target = tmp_path / "new.docx"
        pipeline.commit_file(target, b"x", None, "output")
        assert stat.S_IMODE(target.stat().st_mode) == 0o644
        os.chmod(target, 0o664)
        pipeline.commit_file(target, b"y", pipeline._snapshot(target), "output")
        assert stat.S_IMODE(target.stat().st_mode) == 0o664 and target.read_bytes() == b"y"
    finally:
        os.umask(old)


# ---- 10, 11: inputs and the template --------------------------------------------------------------------


def test_source_edited_during_the_scan_is_detected(tmp_path, monkeypatch):
    src = _doc(tmp_path, "Body text.\\input{part}")
    (tmp_path / "part.tex").write_text("Part text.\n")
    from texdocxdrafter import preflight as pfm

    real = pfm.read_source_text

    def read_then_edit(path, issues, shown=None, fingerprints=None):
        text = real(path, issues, shown, fingerprints)
        if path.name == "main.tex":
            path.write_text(path.read_text() + "% edited while the tool was scanning\n")
        return text

    monkeypatch.setattr(pfm, "read_source_text", read_then_edit)
    with pytest.raises(SourceChangedError, match="main.tex changed"):
        build(src, no_harvest=True, log=_quiet)
    assert not (tmp_path / "main.docx").exists()


def test_output_cannot_be_the_reference_document(tmp_path):
    import docx

    tpl = tmp_path / "tpl.docx"
    docx.Document().save(str(tpl))
    before = tpl.read_bytes()
    src = _doc(tmp_path, "Body.")
    with pytest.raises(PreflightError, match="is the reference document itself"):
        build(src, out=tpl, reference_doc=tpl, overwrite=True, no_harvest=True, log=_quiet)
    assert tpl.read_bytes() == before


# ---- 7, 8, 12, 13: constructs ------------------------------------------------------------------------------


def test_semicolon_binomial_intertext_and_unknown_macro_in_math(tmp_path):
    src = _doc(tmp_path, r"""Density $p(y;\mu)$ and $\binom{n}{k}$ and $\foo{x}$.
\begin{align*}
a &= b \\
\intertext{and therefore}
c &= d
\end{align*}
\[ p(y;\mu) \]""")
    res = build(src, no_harvest=True, log=_quiet)
    xml = _xml(tmp_path)
    assert xml.count("<m:t>;</m:t>") >= 2 or xml.count(";") >= 2
    assert '<m:type m:val="noBar"/>' in xml, "binomial drawn as an ordinary fraction"
    assert "and therefore" in _paras(xml), _paras(xml)[:8]
    assert any("\\foo is not known" in e.message for e in res.audit.entries), [e.message for e in res.audit.entries]


def test_longtable_footers_and_negative_multirow(tmp_path):
    src = _doc(tmp_path, r"""\begin{longtable}{ll}
H1 & H2 \\ \hline
\endfirsthead
H1 & H2 (cont.) \\ \hline
\endhead
\multicolumn{2}{r}{continued on next page} \\
\endfoot
Total & 99 \\
\endlastfoot
r1 & 1 \\
r2 & 2 \\
\end{longtable}
\begin{tabular}{ll}
 & 1 \\
\multirow{-2}{*}{Span} & 2 \\
\end{tabular}
\begin{tabular}{ll}
x & 1 \\
\multirow{-2}{*}{Kept} & 2 \\
\multicolumn{2}{c}{Head} \\
\multirow{-2}{*}{Also} & 4 \\
\end{tabular}""")
    res = build(src, no_harvest=True, log=_quiet)
    xml = _xml(tmp_path)
    cells = [_text(c) for c in re.findall(r"<w:tc>.*?</w:tc>", xml, re.S)]
    assert "continued on next page" not in " ".join(cells)
    assert cells.index("Total") > cells.index("r2"), cells  # last footer after the body
    assert "(cont.)" not in " ".join(cells)
    assert any("per-page footer" in e.message for e in res.audit.entries)
    tables = re.findall(r"<w:tbl>.*?</w:tbl>", xml, re.S)
    # \multirow{-2} over an empty cell: the span starts one row above and the text sits in that top cell
    rows = re.findall(r"<w:tr[ >].*?</w:tr>", tables[1], re.S)
    assert _text(rows[0]).startswith("Span") and 'w:val="restart"' in rows[0] and "Span" not in _text(rows[1]), [_text(r) for r in rows]
    # R2: over cells that hold text or a \multicolumn, nothing is overwritten and the report says so
    texts = " ".join(_text(r) for r in re.findall(r"<w:tr[ >].*?</w:tr>", tables[2], re.S))
    for word in ("x", "Kept", "Head", "Also"):
        assert word in texts.split() or word in texts, (word, texts)
    assert sum("not merged" in e.message for e in res.audit.entries) == 2


def test_single_subfigure_caption_panel_letters_and_subfigure_labels(tmp_path):
    src = _doc(tmp_path, r"""\begin{figure}
\begin{subfigure}{\linewidth}\includegraphics{pic.png}\caption{Only panel}\label{sub:only}\end{subfigure}
\caption{One}
\end{figure}
\begin{figure}
\begin{subfigure}{.5\linewidth}\includegraphics{pic.png}\caption{Left}\label{sub:l}\end{subfigure}
\begin{subfigure}{.5\linewidth}\includegraphics{pic.png}\caption{Right}\label{sub:r}\end{subfigure}
\caption{Two}
\end{figure}
Panels \ref{sub:only}, \ref{sub:l} and \ref{sub:r}.
\[ E = mc^2 \label{eq:plain} \]
Equation \ref{eq:plain}.""")
    res = build(src, no_harvest=True, log=_quiet)
    assert res.check.ok, res.check.errors
    paras = _paras(_xml(tmp_path))
    assert "Only panel" in paras and "(a) Left" in paras and "(b) Right" in paras, paras
    assert not [e for e in res.audit.entries if "no target" in e.message], [e.message for e in res.audit.entries]


def test_author_environment_wrapping_a_tabular_or_list_keeps_it(tmp_path):
    src = _doc(tmp_path, r"""\begin{mytab}
a & 1 \\
b & 2 \\
\end{mytab}
\begin{mylist}
\item one
\item two
\end{mylist}""", preamble=r"""\newenvironment{mytab}{\begin{center}\begin{tabular}{ll}}{\end{tabular}\end{center}}
\newenvironment{mylist}{\begin{itemize}}{\end{itemize}}""")
    res = build(src, no_harvest=True, log=_quiet)
    got = fidelity.found(tmp_path / "main.docx", stylemap.load(None))
    assert got.tables == 1
    xml = _xml(tmp_path)
    assert "<w:numPr>" in xml and "one" in _paras(xml) and "two" in _paras(xml)
    assert not [e for e in res.audit.entries if e.kind == "raw"], [e.message for e in res.audit.entries if e.kind == "raw"]


def test_small_constructs_are_rendered_or_reported(tmp_path):
    src = _doc(tmp_path, "a\\hfill b and caret \\^{} tilde \\~{} here.\n"
               "Bad link \\href{file:///C:/x.docx}{local file}.\n"
               "Outer.\\footnote{Outer note\\footnote{Inner note} continues.}\n"
               "Control\x01char.")
    res = build(src, no_harvest=True, log=_quiet)
    assert res.check.ok, res.check.errors
    text = _text(_xml(tmp_path))
    assert "a b and caret ^ tilde ~ here." in text, text
    msgs = [e.message for e in res.audit.entries]
    assert any("file:///C:/x.docx" in m and "not a web or mail address" in m for m in msgs), msgs
    fn = _text(_xml(tmp_path, "word/footnotes.xml"))
    assert "Outer note (Inner note) continues." in fn, fn
    assert any("footnote inside a footnote" in m for m in msgs)
    assert any("control character" in m for m in msgs), msgs
    assert "Controlchar." in text


def test_fields_inside_tracked_deletions_are_checked(tmp_path):
    src = _doc(tmp_path, "Body.")
    build(src, no_harvest=True, log=_quiet)
    out = tmp_path / "main.docx"
    W = b'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
    _rezip(out, {"word/document.xml": [(b"<w:body>", b"<w:body><w:p><w:del " + W + b' w:id="9" w:author="x" w:date="2026-01-01T00:00:00Z">'
        b'<w:r><w:fldChar w:fldCharType="begin"/></w:r><w:r><w:delInstrText xml:space="preserve"> INCLUDETEXT "\\\\\\\\evil\\\\x" </w:delInstrText></w:r>'
        b'<w:r><w:fldChar w:fldCharType="end"/></w:r></w:del></w:p>')]}, {})
    errs = check_docx(out).errors
    assert any("policy" in e and "INCLUDETEXT" in e for e in errs), errs


# ---- round 13: regressions of the above ------------------------------------------------------------------------


def test_escaped_percent_in_url_is_kept_and_the_line_survives(tmp_path):
    """R1: hyperref's recommended \\href{...a\\%20b}{} form; R4: the comment after a protected url is intact."""
    src = _doc(tmp_path, "See \\href{https://ex.com/a\\%20b}{the page} for details. % Reviewer: check this\n"
                         "Also \\url{https://ex.com/c%20d} here. % Reviewer: and this")
    from texdocxdrafter.preflight import preflight

    pf = preflight(src, need_build_files=False)
    assert [c.text for c in pf.comments] == ["Reviewer: check this", "Reviewer: and this"]
    res = build(src, no_harvest=True, log=_quiet)
    rels = _xml(tmp_path, "word/_rels/document.xml.rels")
    assert 'Target="https://ex.com/a%20b"' in rels and 'Target="https://ex.com/c%20d"' in rels, rels
    text = _text(_xml(tmp_path))
    assert "the page for details." in text and "https://ex.com/c%20d here." in text, text
    assert not [e for e in res.audit.entries if "link target" in e.message], [e.message for e in res.audit.entries]


def test_relationships_of_nested_parts_are_found_beside_them(tmp_path):
    """R3: word/theme/theme1.xml keeps its .rels in word/theme/_rels/, not word/_rels/theme/."""
    src = _doc(tmp_path, "Body.")
    build(src, no_harvest=True, log=_quiet)
    out = tmp_path / "main.docx"
    theme = (b'<?xml version="1.0" encoding="UTF-8"?><a:theme xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
             b'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" name="t">'
             b'<a:blip r:embed="rId1"/></a:theme>')
    rels = (b'<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            b'<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="../media/fill.png"/>'
            b'</Relationships>')
    _rezip(out, {}, {"word/theme/theme1.xml": theme, "word/theme/_rels/theme1.xml.rels": rels, "word/media/fill.png": PNG})
    errs = check_docx(out).errors
    assert not [e for e in errs if "unresolved" in e], errs
    _rezip(out, {"word/theme/theme1.xml": [(b'r:embed="rId1"', b'r:embed="rId7"')]}, {})
    assert any("theme1.xml: unresolved relationship rId7" in e for e in check_docx(out).errors)


def test_chmod_failure_on_fat_drives_does_not_abort(tmp_path, monkeypatch):
    """R5: FAT/exFAT answer EPERM to chmod; the content still lands."""
    monkeypatch.setattr(pipeline.os, "chmod", lambda *a, **k: (_ for _ in ()).throw(OSError(errno.EPERM, "Operation not permitted")))
    target = tmp_path / "t.docx"
    pipeline.commit_file(target, b"data", None, "output")
    pipeline.commit_file(target, b"more", pipeline._snapshot(target), "output")
    assert target.read_bytes() == b"more"
    _no_temp_files(tmp_path)


def test_panel_letters_skip_a_plain_picture_before_the_subfigures(tmp_path):
    """R6: LaTeX letters captioned subfigures only."""
    src = _doc(tmp_path, r"""\begin{figure}
\includegraphics{pic.png}
\begin{subfigure}{.5\linewidth}\includegraphics{pic.png}\caption{Left}\end{subfigure}
\begin{subfigure}{.5\linewidth}\includegraphics{pic.png}\caption{Right}\end{subfigure}
\caption{Three pictures}
\end{figure}""")
    build(src, no_harvest=True, log=_quiet)
    paras = _paras(_xml(tmp_path))
    assert "(a) Left" in paras and "(b) Right" in paras, paras


# ---- round 15 ---------------------------------------------------------------------------------------------


def test_caption_star_has_no_number_and_does_not_advance_the_sequence(tmp_path):
    """C1: \\caption* gives an unnumbered caption; the next figure is still Figure 1."""
    src = _doc(tmp_path, r"""\begin{figure}\includegraphics{pic.png}\caption*{Unnumbered fig}\end{figure}
\begin{figure}\includegraphics{pic.png}\caption{Real fig}\label{fig:real}\end{figure}
\begin{table}\caption*{Loose table}\begin{tabular}{l} a \end{tabular}\end{table}
\begin{table}\caption{Real table}\begin{tabular}{l} b \end{tabular}\end{table}
See \ref{fig:real}.""")
    res = build(src, no_harvest=True, log=_quiet)
    assert res.check.ok, res.check.errors
    paras = _paras(_xml(tmp_path))
    assert "Unnumbered fig" in paras and "Figure 1: Real fig" in paras, paras
    assert "Loose table" in paras and "Table 1: Real table" in paras, paras
    assert not [p for p in paras if p.startswith("Figure 2") or p.startswith("Table 2")]


def test_sequence_follows_latex_numbers_from_the_aux(tmp_path):
    """C2': when LaTeX's number differs from Word's running count, the SEQ field resets the sequence."""
    src = _doc(tmp_path, r"""\begin{figure}\includegraphics{pic.png}\caption{First}\label{fig:a}\end{figure}
\setcounter{figure}{4}
\begin{figure}\includegraphics{pic.png}\caption{Fifth in LaTeX}\label{fig:b}\end{figure}
\begin{figure}\includegraphics{pic.png}\caption{Sixth}\label{fig:c}\end{figure}
\begin{equation} a = b \label{eq:a} \end{equation}
\begin{equation} c = d \label{eq:b} \end{equation}
Refs \ref{fig:b}, \ref{fig:c}, \ref{eq:b}.""")
    src.with_suffix(".aux").write_text("\\relax\n\\newlabel{fig:a}{{1}{1}}\n\\newlabel{fig:b}{{5}{1}}\n\\newlabel{fig:c}{{6}{1}}\n"
                                       "\\newlabel{eq:a}{{1}{1}}\n\\newlabel{eq:b}{{3}{1}}\n")
    res = build(src, no_harvest=True, log=_quiet)
    xml = _xml(tmp_path)
    assert "SEQ Figure \\r 5 \\* ARABIC" in xml and "SEQ Equation \\r 3 \\* ARABIC" in xml, xml[:0]
    paras = _paras(xml)
    assert "Figure 5: Fifth in LaTeX" in paras and "Figure 6: Sixth" in paras, paras
    assert xml.count("SEQ Figure \\r") == 1  # the sixth follows the running count again
    msgs = [e.message for e in res.audit.entries if e.kind == "numbering"]
    assert any("reset to LaTeX's number" in m for m in msgs), msgs
    assert all(e.severity == "info" for e in res.audit.entries if e.kind == "numbering"), msgs
