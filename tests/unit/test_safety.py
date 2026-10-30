"""Safety properties: sandboxed compiles, private work directories, output vetting, include limits, comments opt-in,
image cache keys and tool limits."""
import os
import shutil
import time
import zipfile
from pathlib import Path

import pytest

from texdocxdrafter import sandbox
from texdocxdrafter import pipeline
from texdocxdrafter.pipeline import WORK_MARKER, build, output_path, write_atomic
from texdocxdrafter.auxreader import read_aux
from texdocxdrafter.preflight import PreflightError, preflight, side_preamble
from texdocxdrafter.writer.images import find_graphic, to_raster

HAVE_PDFLATEX = shutil.which("pdflatex") is not None
HAVE_LUALATEX = shutil.which("lualatex") is not None
HAVE_POPPLER = shutil.which("pdftoppm") is not None and shutil.which("pdfinfo") is not None

DOC = r"""\documentclass{article}
\begin{document}
Hello world. % a private note
\end{document}
"""


def _write(tmp_path, text=DOC, name="main.tex", compiled=True):
    src = tmp_path / name
    src.write_text(text)
    if compiled:
        src.with_suffix(".aux").write_text("\\relax\n")
    return src


def _compile(src: Path):
    r = sandbox.run(["pdflatex", "-interaction=batchmode", src.name], cwd=src.parent, timeout=60)
    assert src.with_suffix(".pdf").exists(), r.stderr


# ---- 1. TeX side compiles cannot write outside their directory or run shell commands --------------------

@pytest.mark.skipif(not HAVE_PDFLATEX, reason="pdflatex not installed")
def test_tex_compile_cannot_escape(tmp_path):
    victim = tmp_path / "victim.txt"
    victim.write_text("untouched")
    shell_mark = tmp_path / "shell-ran.txt"
    work = tmp_path / "work"
    work.mkdir()
    tex = work / "evil.tex"
    tex.write_text(r"""\documentclass{article}
\begin{document}
\newwrite\w
\immediate\openout\w=%s
\immediate\write\w{pwned}
\immediate\closeout\w
\immediate\write18{touch %s}
x
\end{document}
""" % (victim, shell_mark))
    r = sandbox.compile_tex("pdflatex", work, "evil", tex, [work], timeout=60, check_policy=False)
    assert r.returncode != -1
    assert victim.read_text() == "untouched"
    assert not shell_mark.exists()
    assert (work / "evil.log").exists()  # the engine ran and wrote only into its work directory


def test_timeout_is_enforced():
    t0 = time.time()
    r = sandbox.run(["sleep", "5"], cwd="/tmp", timeout=1)
    assert r.returncode == -1 and "timed out" in r.stderr
    assert time.time() - t0 < 4


# ---- 2. Cleanup never touches files the tool did not create ------------------------------------------

@pytest.mark.skipif(not HAVE_PDFLATEX, reason="pdflatex not installed")
def test_config_dir_contents_survive_and_work_is_kept_separately(tmp_path):
    src = _write(tmp_path)
    cfg = tmp_path / "main.texdocxdrafter"
    cfg.mkdir()
    decoy = cfg / "my-notes.txt"
    decoy.write_text("mine")
    foreign = cfg / "work"
    foreign.mkdir()
    (foreign / "precious.txt").write_text("not ours")
    build(src, out=tmp_path / "o.docx", no_harvest=True, keep_work=True, log=lambda *_: None)
    assert decoy.read_text() == "mine"
    assert (foreign / "precious.txt").read_text() == "not ours"  # a foreign 'work' dir is left alone
    kept = [d for d in cfg.iterdir() if d.is_dir() and d.name.startswith("work-")]
    assert len(kept) == 1 and (kept[0] / WORK_MARKER).exists() and (kept[0] / "tree.json").exists()
    # a second run replaces only the directory carrying our marker
    build(src, out=tmp_path / "o.docx", no_harvest=True, keep_work=True, log=lambda *_: None)
    assert (foreign / "precious.txt").exists()


def test_no_leftover_temp_dirs(tmp_path):
    src = _write(tmp_path)
    before = {p.name for p in Path("/tmp").glob("texdocxdrafter-*")}
    build(src, out=tmp_path / "o.docx", no_harvest=True, log=lambda *_: None)
    after = {p.name for p in Path("/tmp").glob("texdocxdrafter-*")}
    assert after <= before


# ---- 3. Output never overwrites sources or foreign documents -----------------------------------------

def test_output_cannot_be_a_source_or_non_docx(tmp_path):
    src = _write(tmp_path)
    pf = preflight(src, need_build_files=False)
    with pytest.raises(PreflightError):
        output_path(pf, src)
    with pytest.raises(PreflightError):
        output_path(pf, tmp_path / "main.tex")
    with pytest.raises(PreflightError):
        output_path(pf, tmp_path / "out.txt")


def test_foreign_docx_is_protected_but_own_output_is_replaced(tmp_path):
    src = _write(tmp_path)
    out = tmp_path / "main.docx"
    out.write_bytes(b"a coauthor's draft, not a zip at all")
    with pytest.raises(PreflightError):
        build(src, no_harvest=True, log=lambda *_: None)
    assert out.read_bytes().startswith(b"a coauthor")
    build(src, no_harvest=True, overwrite=True, log=lambda *_: None)
    assert zipfile.is_zipfile(out)
    # our own untouched output may be replaced without the flag
    build(src, no_harvest=True, log=lambda *_: None)
    # a document saved by someone else (lastModifiedBy changed) is protected again
    edited = tmp_path / "edited.docx"
    with zipfile.ZipFile(out) as zin, zipfile.ZipFile(edited, "w") as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "docProps/core.xml":
                data = data.replace(b"<cp:lastModifiedBy>texdocxdrafter</cp:lastModifiedBy>",
                                    b"<cp:lastModifiedBy>A Coauthor</cp:lastModifiedBy>")
            zout.writestr(item, data)
    with pytest.raises(PreflightError):
        build(src, out=edited, no_harvest=True, log=lambda *_: None)


# ---- 4. \input and graphics stay inside allowed directories -----------------------------------------

def test_input_outside_source_dir_is_refused_unless_allowed(tmp_path):
    outside = tmp_path / "private.txt"
    outside.write_text("SECRET")
    proj = tmp_path / "proj"
    proj.mkdir()
    src = _write(proj, text="\\documentclass{article}\\begin{document}\\input{../private.txt}\\end{document}\n")
    pf = preflight(src, need_build_files=False)
    assert "SECRET" not in pf.body
    assert any("outside the source directory" in i for i in pf.issues)
    pf = preflight(src, need_build_files=False, allow_dirs=[tmp_path])
    assert "SECRET" in pf.body


def test_symlinked_input_is_resolved(tmp_path):
    outside = tmp_path / "private.tex"
    outside.write_text("SECRET")
    proj = tmp_path / "proj"
    proj.mkdir()
    os.symlink(outside, proj / "link.tex")
    src = _write(proj, text="\\documentclass{article}\\begin{document}\\input{link}\\end{document}\n")
    pf = preflight(src, need_build_files=False)
    assert "SECRET" not in pf.body


def test_graphics_outside_source_dir_not_found(tmp_path):
    (tmp_path / "fig.png").write_bytes(b"\x89PNG")
    proj = tmp_path / "proj"
    proj.mkdir()
    assert find_graphic("../fig.png", proj, []) is None
    assert find_graphic("../fig.png", proj, [], [proj, tmp_path]).resolve() == (tmp_path / "fig.png").resolve()
    assert find_graphic("fig", tmp_path, []) == tmp_path / "fig.png"


# ---- 5. Comments are exported only on request -----------------------------------------------------

def test_comments_are_opt_in(tmp_path):
    src = _write(tmp_path)
    out = tmp_path / "o.docx"
    build(src, out=out, no_harvest=True, log=lambda *_: None)
    assert "word/comments.xml" not in zipfile.ZipFile(out).namelist()
    build(src, out=out, no_harvest=True, comments=True, log=lambda *_: None)
    assert "private note" in zipfile.ZipFile(out).read("word/comments.xml").decode()


# ---- 6. Different figures never share a raster ---------------------------------------------------

@pytest.mark.skipif(not (HAVE_PDFLATEX and HAVE_POPPLER), reason="pdflatex or poppler not installed")
def test_distinct_figures_get_distinct_rasters(tmp_path):
    pngs = []
    for d, txt in (("a", "AAAA"), ("b", "BBBBBBBBBBBB")):
        sub = tmp_path / d
        sub.mkdir()
        fig = sub / "figure.tex"
        fig.write_text("\\documentclass{article}\\pagestyle{empty}\\begin{document}%s\\end{document}\n" % txt)
        _compile(fig)
        pngs.append(to_raster(sub / "figure.pdf", tmp_path / "work"))
    assert pngs[0] != pngs[1]
    assert pngs[0].read_bytes() != pngs[1].read_bytes()


# ---- 7. Converters run with limits ------------------------------------------------------------------

@pytest.mark.skipif(not (HAVE_PDFLATEX and HAVE_POPPLER), reason="pdflatex or poppler not installed")
def test_huge_page_is_rasterised_within_pixel_limit(tmp_path):
    fig = tmp_path / "big.tex"
    fig.write_text("\\documentclass{article}\\usepackage[paperwidth=200in,paperheight=200in]{geometry}"
                   "\\pagestyle{empty}\\begin{document}x\\end{document}\n")
    _compile(fig)
    png = to_raster(tmp_path / "big.pdf", tmp_path / "work")
    assert png is not None
    from docx.image.image import Image

    im = Image.from_file(str(png))
    assert max(im.px_width, im.px_height) <= 6000


def test_side_compile_preamble_rebases_allowed_relative_paths(tmp_path):
    shared = tmp_path / "shared"
    shared.mkdir()
    (shared / "pre.tex").write_text("\\usepackage{xcolor}\n")
    (shared / "refs.bib").write_text("")
    proj = tmp_path / "proj"
    proj.mkdir()
    src = _write(proj, text="\\documentclass{article}\n\\input{../shared/pre}\n\\addbibresource{../shared/refs.bib}\n"
                            "\\graphicspath{{../shared/}{figs/}}\n\\begin{document}x\\end{document}\n")
    pf = preflight(src, need_build_files=False)
    assert "../shared/pre" in side_preamble(pf)  # not allowed: left as written (fails loudly in the side compile)
    pf = preflight(src, need_build_files=False, allow_dirs=[shared])
    side = side_preamble(pf)
    assert "\\input{%s}" % (shared / "pre.tex") in side
    assert "\\addbibresource{%s}" % (shared / "refs.bib") in side
    assert "\\graphicspath{{%s/}{figs/}}" % shared in side


LUA_EVIL = r"""\documentclass{article}
\begin{document}
\directlua{
  local f = io.open("%(victim)s", "w") if f then f:write("pwned") f:close() end
  local g = io.open("%(new)s", "w") if g then g:write("pwned") g:close() end
  os.remove("%(victim)s")
  os.execute("touch %(shell)s")
  local p = io.popen("touch %(popen)s") if p then p:close() end
  local ok, lfs = pcall(require, "lfs") if ok then lfs.mkdir("%(mkdir)s") end
}
x
\end{document}
"""


def _lua_escape_attempt(tmp_path):
    victim = tmp_path / "victim.txt"
    victim.write_text("untouched")
    names = {k: str(tmp_path / f"{k}.txt") for k in ("new", "shell", "popen", "mkdir")}
    names["victim"] = str(victim)
    work = tmp_path / "work"
    work.mkdir()
    tex = work / "evil.tex"
    tex.write_text(LUA_EVIL % names)
    r = sandbox.compile_tex("lualatex", work, "evil", tex, [work], timeout=120, check_policy=False)
    assert r.returncode != -1, r.stderr
    assert (work / "evil.log").exists()
    assert victim.read_text() == "untouched"
    for k in ("new", "shell", "popen", "mkdir"):
        assert not Path(names[k]).exists(), k


@pytest.mark.skipif(not HAVE_LUALATEX, reason="lualatex not installed")
def test_luatex_cannot_escape_with_os_sandbox(tmp_path):
    if not sandbox.have_os_sandbox():
        pytest.skip("bubblewrap not available")
    _lua_escape_attempt(tmp_path)


@pytest.mark.skipif(not HAVE_LUALATEX, reason="lualatex not installed")
def test_luatex_cannot_escape_without_os_sandbox(tmp_path, monkeypatch):
    monkeypatch.setitem(sandbox._disabled, "bwrap", True)
    assert not sandbox.have_os_sandbox()
    _lua_escape_attempt(tmp_path)


@pytest.mark.skipif(not HAVE_PDFLATEX, reason="pdflatex not installed")
def test_pdftex_cannot_escape_without_os_sandbox(tmp_path, monkeypatch):
    monkeypatch.setitem(sandbox._disabled, "bwrap", True)
    test_tex_compile_cannot_escape(tmp_path)


# ---- second review round -------------------------------------------------------------------------

LUA_CHDIR = r"""\documentclass{article}
\begin{document}
\directlua{
  local lfs = require("lfs")
  lfs.chdir("..")
  local f = io.open("victim.txt", "w") if f then f:write("pwned") f:close() end
  local g = io.open("../victim.txt", "w") if g then g:write("pwned") g:close() end
  lfs.link("work/evil.tex", "%(link)s")
  lfs.link("work/evil.tex", "%(symlink)s", true)
  local s = io.open("%(secret)s", "r")
  if s then texio.write_nl("SECRET-CONTENT:" .. (s:read("a") or "")) s:close() end
}
x
\end{document}
"""


def _lua_chdir_attempt(tmp_path):
    victim = tmp_path / "victim.txt"
    victim.write_text("untouched")
    secret = tmp_path / "secret.txt"
    secret.write_text("TOPSECRET")
    work = tmp_path / "work"
    work.mkdir()
    names = {"link": str(tmp_path / "link.txt"), "symlink": str(tmp_path / "symlink.txt"), "secret": str(secret)}
    tex = work / "evil.tex"
    tex.write_text(LUA_CHDIR % names)
    r = sandbox.compile_tex("lualatex", work, "evil", tex, [work], timeout=120, check_policy=False)
    assert r.returncode != -1, r.stderr
    assert victim.read_text() == "untouched"
    assert not (tmp_path / "link.txt").exists() and not (tmp_path / "symlink.txt").is_symlink()
    return (work / "evil.log").read_text(errors="replace")


@pytest.mark.skipif(not HAVE_LUALATEX, reason="lualatex not installed")
def test_luatex_chdir_and_link_bypass_closed_without_os_sandbox(tmp_path, monkeypatch):
    monkeypatch.setitem(sandbox._disabled, "bwrap", True)
    _lua_chdir_attempt(tmp_path)


@pytest.mark.skipif(not HAVE_LUALATEX, reason="lualatex not installed")
def test_os_sandbox_hides_files_outside_the_inputs(tmp_path):
    if not sandbox.have_os_sandbox():
        pytest.skip("bubblewrap not available")
    log = _lua_chdir_attempt(tmp_path)
    assert "TOPSECRET" not in log  # the secret next to the work directory is not even readable


def test_no_tex_engine_runs_without_os_sandbox(monkeypatch):
    assert sandbox.sandbox_missing() is None or not sandbox.have_os_sandbox()
    monkeypatch.setitem(sandbox._disabled, "bwrap", True)
    assert "bubblewrap" in sandbox.sandbox_missing()


def test_environment_is_filtered():
    env = sandbox.filtered_env({"PATH": "/bin", "AWS_SECRET_ACCESS_KEY": "x", "TEXMFHOME": "/h", "LC_ALL": "C",
                                "HOME": "/home/u", "GITHUB_TOKEN": "y"})
    assert set(env) == {"PATH", "TEXMFHOME", "LC_ALL", "HOME"}


def test_edited_docx_is_protected_even_if_properties_survive(tmp_path):
    src = _write(tmp_path)
    out = tmp_path / "main.docx"
    build(src, no_harvest=True, log=lambda *_: None)
    import docx

    d = docx.Document(str(out))  # a python-docx edit keeps lastModifiedBy; the content hash still changes
    d.add_paragraph("a coauthor's addition")
    d.save(str(out))
    with pytest.raises(PreflightError):
        build(src, no_harvest=True, log=lambda *_: None)
    assert "coauthor" in zipfile.ZipFile(out).read("word/document.xml").decode()


def test_symlinked_output_and_report_are_refused(tmp_path):
    src = _write(tmp_path)
    target = tmp_path / "elsewhere.docx"
    target.write_text("keep me")
    out = tmp_path / "main.docx"
    os.symlink(target, out)
    with pytest.raises(PreflightError):
        build(src, no_harvest=True, overwrite=True, log=lambda *_: None)
    assert target.read_text() == "keep me"
    out.unlink()
    rep_target = tmp_path / "elsewhere.md"
    rep_target.write_text("keep me too")
    os.symlink(rep_target, tmp_path / "main.texdocxdrafter-report.md")
    with pytest.raises(PreflightError):
        build(src, no_harvest=True, overwrite=True, log=lambda *_: None)
    assert rep_target.read_text() == "keep me too"


def test_write_atomic_never_follows_a_planted_symlink(tmp_path):
    target = tmp_path / "victim.txt"
    target.write_text("untouched")
    # even a temporary name already present as a symlink is not written through: mkstemp creates exclusively
    for name in (".out.docx.tmp", ".out.docx.abc.tmp"):
        os.symlink(target, tmp_path / name)
    write_atomic(tmp_path / "out.docx", b"data")
    assert target.read_text() == "untouched"
    assert (tmp_path / "out.docx").read_bytes() == b"data"


def test_report_is_replaced_when_untouched_but_protected_when_edited(tmp_path):
    src = _write(tmp_path)
    build(src, out=tmp_path / "o.docx", no_harvest=True, log=lambda *_: None)
    rep = tmp_path / "o.texdocxdrafter-report.md"
    build(src, out=tmp_path / "o.docx", no_harvest=True, log=lambda *_: None)  # untouched report: fine
    rep.write_text(rep.read_text() + "\nmy annotation\n")
    with pytest.raises(PreflightError):
        build(src, out=tmp_path / "o.docx", no_harvest=True, log=lambda *_: None)
    assert "my annotation" in rep.read_text()


@pytest.mark.skipif(not (HAVE_PDFLATEX and HAVE_POPPLER), reason="pdflatex or poppler not installed")
def test_figure_names_with_dots_convert(tmp_path):
    fig = tmp_path / "figure.v1.tex"
    fig.write_text("\\documentclass{article}\\pagestyle{empty}\\begin{document}dots\\end{document}\n")
    _compile(fig)
    png = to_raster(tmp_path / "figure.v1.pdf", tmp_path / "work")
    assert png is not None and png.exists() and png.suffix == ".png"


def test_timeout_partial_output_is_text():
    r = sandbox.run(["sh", "-c", "echo partial; sleep 5"], cwd="/tmp", timeout=1)
    assert r.returncode == -1 and isinstance(r.stdout, str) and "partial" in r.stdout


def test_aux_self_include_terminates_and_outside_aux_is_ignored(tmp_path):
    (tmp_path / "sub").mkdir()
    main = tmp_path / "main.aux"
    main.write_text("\\newlabel{a}{{1}{1}}\n\\@input{sub/../main.aux}\n\\@input{../outside.aux}\n\\@input{sub/x.aux}\n")
    (tmp_path / "sub" / "x.aux").write_text("\\newlabel{b}{{2}{1}}\n\\@input{../main.aux}\n")
    (tmp_path.parent / "outside.aux").write_text("\\newlabel{c}{{3}{1}}\n")
    try:
        labels = read_aux(main, [tmp_path])
    finally:
        (tmp_path.parent / "outside.aux").unlink()
    assert set(labels) == {"a", "b"}


# ---- third review round --------------------------------------------------------------------------

def test_compile_tex_is_the_policy_boundary(tmp_path, monkeypatch):
    """Every side compile goes through compile_tex, which refuses to start an engine without an OS sandbox."""
    monkeypatch.setitem(sandbox._disabled, "bwrap", True)
    work = tmp_path / "work"
    work.mkdir()
    tex = work / "x.tex"
    tex.write_text(DOC)
    for engine in ("pdflatex", "xelatex", "lualatex"):
        with pytest.raises(sandbox.SandboxPolicyError):
            sandbox.compile_tex(engine, work, "x", tex, [work], timeout=30)
    assert not (work / "x.log").exists()  # no engine ran


def test_harvest_command_without_os_sandbox_writes_the_side_document(tmp_path, monkeypatch, capsys):
    from texdocxdrafter import cli

    monkeypatch.setitem(sandbox._disabled, "bwrap", True)
    src = _write(tmp_path)
    assert cli.main(["harvest", str(src)]) == 1
    err = capsys.readouterr().err
    assert "harvest.tex" not in err or "compile" in err
    assert (tmp_path / "main.texdocxdrafter" / "harvest.tex").is_file()
    assert "pdflatex harvest" in err
    assert not list(tmp_path.glob("**/harvest.log"))  # the tool itself ran nothing


@pytest.mark.skipif(not HAVE_PDFLATEX, reason="pdflatex not installed")
def test_pdftex_cannot_read_outside_the_inputs_with_os_sandbox(tmp_path):
    if not sandbox.have_os_sandbox():
        pytest.skip("bubblewrap not available")
    secret = tmp_path / "secret.tex"
    secret.write_text(r"\message{TOPSECRET}")
    work = tmp_path / "work"
    work.mkdir()
    tex = work / "peek.tex"
    tex.write_text("\\documentclass{article}\\begin{document}\\input{%s} x\\end{document}\n" % secret)
    r = sandbox.compile_tex("pdflatex", work, "peek", tex, [work], timeout=60)
    assert r.returncode != -1
    assert "TOPSECRET" not in (work / "peek.log").read_text(errors="replace")


def _build_with_intervening(src, tmp_path, during, **kw):
    """Run build(); ``during`` runs after the targets were vetted, before anything is written."""
    real = pipeline.parse_document

    def hooked(*a, **k):
        during()
        return real(*a, **k)

    pipeline.parse_document = hooked
    try:
        return build(src, no_harvest=True, log=lambda *_: None, **kw)
    finally:
        pipeline.parse_document = real


def test_file_created_during_the_run_is_not_overwritten(tmp_path):
    src = _write(tmp_path)
    out = tmp_path / "main.docx"
    with pytest.raises(PreflightError):
        _build_with_intervening(src, tmp_path, lambda: out.write_text("a coauthor saved this meanwhile"))
    assert out.read_text() == "a coauthor saved this meanwhile"
    assert not list(tmp_path.glob(".main.docx.*.tmp"))


def test_file_changed_during_the_run_is_not_overwritten(tmp_path):
    src = _write(tmp_path)
    out = tmp_path / "main.docx"
    build(src, no_harvest=True, log=lambda *_: None)  # our own output: a rerun would normally replace it
    with pytest.raises(PreflightError):
        _build_with_intervening(src, tmp_path, lambda: out.write_text("edited meanwhile"))
    assert out.read_text() == "edited meanwhile"
    # even with --overwrite, a file that changed between vetting and commit is left alone
    with pytest.raises(PreflightError):
        _build_with_intervening(src, tmp_path, lambda: out.write_text("edited again"), overwrite=True)
    assert out.read_text() == "edited again"


def test_registry_never_replaces_a_source_or_foreign_file(tmp_path):
    cfg = tmp_path / "main.texdocxdrafter"
    cfg.mkdir()
    reg = cfg / "outputs.json"
    # 1. the registry path is pulled in as a source file
    reg.write_text("included text\n")
    src = _write(tmp_path, DOC.replace("Hello world.", r"\input{main.texdocxdrafter/outputs.json}"))
    with pytest.raises(PreflightError, match="source"):
        build(src, no_harvest=True, log=lambda *_: None)
    assert reg.read_text() == "included text\n"
    # 2. something that is not a hash registry sits at the path
    src = _write(tmp_path)
    for junk in ('{"a": 1}', "not json", "[]"):
        reg.write_text(junk)
        with pytest.raises(PreflightError, match="registry"):
            build(src, no_harvest=True, log=lambda *_: None)
        assert reg.read_text() == junk
    # 3. a registry changed by another run between vetting and commit is left alone
    reg.unlink()
    build(src, no_harvest=True, log=lambda *_: None)
    good = reg.read_text()
    with pytest.raises(PreflightError):
        _build_with_intervening(src, tmp_path, lambda: reg.write_text("{}"))
    assert reg.read_text() == "{}"
    assert good != "{}"


def test_reference_doc_comments_do_not_leak(tmp_path):
    import docx

    tpl = tmp_path / "template.docx"
    d = docx.Document()
    para = d.add_paragraph("template body text")
    d.add_comment(para.runs[0], text="PRIVATE-TEMPLATE-NOTE", author="Reviewer")
    d.save(str(tpl))
    assert "word/comments.xml" in zipfile.ZipFile(tpl).namelist()

    src = _write(tmp_path)
    out = tmp_path / "main.docx"
    build(src, reference_doc=tpl, no_harvest=True, log=lambda *_: None)
    z = zipfile.ZipFile(out)
    assert "word/comments.xml" not in z.namelist()
    assert all(b"PRIVATE-TEMPLATE-NOTE" not in z.read(n) for n in z.namelist())
    assert all(b"Reviewer" not in z.read(n) for n in z.namelist() if n.startswith("word/"))

    build(src, reference_doc=tpl, no_harvest=True, comments=True, log=lambda *_: None)
    z = zipfile.ZipFile(out)
    comments = z.read("word/comments.xml")
    assert b"a private note" in comments and b"PRIVATE-TEMPLATE-NOTE" not in comments


# ---- side compiles done by the user, and bibtex without any compile -------------------------------

BIBLATEX_DOC = r"""\documentclass{article}
\usepackage[style=authoryear,backend=biber]{biblatex}
\usepackage{tikz}
\addbibresource{refs.bib}
\begin{document}
Text \parencite{knuth84} here.
\begin{figure}\centering\begin{tikzpicture}\draw (0,0) -- (1,1);\end{tikzpicture}\caption{A line.}\end{figure}
\printbibliography
\end{document}
"""
BIB = r"""@book{knuth84, author={Donald E. Knuth}, title={The TeXbook}, year={1984}, publisher={Addison-Wesley}}
"""


def _docx_xml(path: Path) -> str:
    return zipfile.ZipFile(path).read("word/document.xml").decode()


@pytest.mark.skipif(not (HAVE_PDFLATEX and HAVE_POPPLER and shutil.which("biber")), reason="pdflatex, biber or poppler missing")
def test_without_os_sandbox_the_user_compiles_the_side_documents(tmp_path, monkeypatch):
    import subprocess

    src = tmp_path / "main.tex"
    src.write_text(BIBLATEX_DOC)
    (tmp_path / "refs.bib").write_text(BIB)
    # the user's own compile of the paper (this is what the tool assumes exists)
    for cmd in (["pdflatex", "-interaction=batchmode", "main"], ["biber", "main"], ["pdflatex", "-interaction=batchmode", "main"]):
        subprocess.run(cmd, cwd=tmp_path, capture_output=True, timeout=120)
    assert (tmp_path / "main.bbl").is_file()

    monkeypatch.setitem(sandbox._disabled, "bwrap", True)
    engine_calls = []
    real_run = sandbox.run

    def spy(cmd, *a, **k):
        if Path(cmd[0]).name in ("pdflatex", "xelatex", "lualatex", "biber", "bibtex"):
            engine_calls.append(cmd[0])
        return real_run(cmd, *a, **k)

    monkeypatch.setattr(sandbox, "run", spy)
    res = build(src, log=lambda *_: None)
    cfg = tmp_path / "main.texdocxdrafter"
    assert (cfg / "harvest.tex").is_file() and (cfg / "tikz.tex").is_file()
    assert not engine_calls  # the tool itself ran no TeX
    msgs = "\n".join(e.message for e in res.audit.entries)
    assert "pdflatex harvest" in msgs and "biber harvest" in msgs and "pdflatex tikz" in msgs
    xml = _docx_xml(tmp_path / "main.docx")
    assert "[knuth84]" in xml and "Knuth" not in xml
    assert "word/media/" not in "\n".join(zipfile.ZipFile(tmp_path / "main.docx").namelist())

    # the user runs the two printed commands
    from texdocxdrafter.sidecompile import compile_command
    from texdocxdrafter.preflight import preflight

    pf = preflight(src)
    for stem, bib in (("harvest", "biber"), ("tikz", None)):
        r = subprocess.run(compile_command(pf, cfg, stem, bib) + " -interaction=batchmode", shell=True,
                           capture_output=True, text=True, timeout=300)
        assert (cfg / f"{stem}.pdf").is_file(), r.stdout[-800:]
    engine_calls.clear()
    res = build(src, log=lambda *_: None)
    assert not engine_calls
    xml = _docx_xml(tmp_path / "main.docx")
    assert "Knuth" in xml and "1984" in xml and "[knuth84]" not in xml
    assert any(n.startswith("word/media/") for n in zipfile.ZipFile(tmp_path / "main.docx").namelist())
    assert res.audit.info["harvest"].startswith("1/1")


NATBIB_DOC = r"""\documentclass{article}
\usepackage[round]{natbib}
\begin{document}
As \citet{smith2020a} showed \citep[see][p.~3]{jones2019,smith2020a}.
\bibliographystyle{plainnat}
\bibliography{refs}
\end{document}
"""
NATBIB_BBL = r"""\begin{thebibliography}{2}
\providecommand{\natexlab}[1]{#1}
\bibitem[Jones and M\"{u}ller(2019)]{jones2019}
K.~Jones and A.~M\"{u}ller.
\newblock \emph{A Book About Stuff}.
\newblock Publisher \& Sons, 2019.

\bibitem[{Smith et~al.(2020{\natexlab{a}})Smith, Jones, and Lee}]{smith2020a}
J.~Smith, K.~Jones, and M.~Lee.
\newblock The title of the paper.
\newblock {\em Journal of Things}, 12(3):45--67, 2020{\natexlab{a}}.

\end{thebibliography}
"""


def test_bibtex_document_needs_no_compile(tmp_path, monkeypatch):
    src = _write(tmp_path, NATBIB_DOC)
    src.with_suffix(".bbl").write_text(NATBIB_BBL)
    monkeypatch.setitem(sandbox._disabled, "bwrap", True)
    monkeypatch.setattr(sandbox, "compile_tex", lambda *a, **k: pytest.fail("a TeX engine was started"))
    res = build(src, log=lambda *_: None)
    assert res.audit.info["harvest source"].startswith("main.bbl")
    assert res.audit.info["harvest"] == "2/2 citations, 2 bibliography entries"
    xml = _docx_xml(tmp_path / "main.docx")
    assert "Smith et al. (2020a)" in xml
    assert "(see Jones and Müller, 2019; Smith et al., 2020a, p. 3)" in xml
    assert "Journal of Things" in xml and "Publisher &amp; Sons" in xml
    import re

    italic = re.findall(r"<w:r>(?:(?!</w:r>).)*(?:<w:i/>|w:val=\"Emphasis\")(?:(?!</w:r>).)*<w:t[^>]*>([^<]*)</w:t>", xml, re.S)
    assert "Journal of Things" in italic and "A Book About Stuff" in italic
    assert not (tmp_path / "main.texdocxdrafter" / "harvest.tex").exists()


# ---- fourth review round --------------------------------------------------------------------------

def _texmfvar() -> Path | None:
    v = sandbox._kpse_var("TEXMFVAR")
    return Path(v) if v and Path(v).is_dir() else None


def test_throwaway_texmfvar_lives_in_the_work_dir(tmp_path):
    env = sandbox.throwaway_texmfvar(tmp_path)
    assert Path(env["TEXMFVAR"]).parent == tmp_path and Path(env["TEXMFVAR"]).is_dir()
    assert env["TEXMFCACHE"].startswith(env["TEXMFVAR"])
    assert not any(str(c) == env["TEXMFVAR"] for c in sandbox.ro_cache_dirs())


@pytest.mark.skipif(not HAVE_LUALATEX, reason="lualatex not installed")
def test_tex_cache_cannot_be_planted_from_the_sandbox(tmp_path):
    """N1: a document must not leave code in the user's TEXMFVAR, where a later ordinary compile would run it."""
    if not sandbox.have_os_sandbox():
        pytest.skip("bubblewrap not available")
    real = _texmfvar()
    if real is None:
        pytest.skip("no TEXMFVAR")
    planted = real / f"texdocxdrafter-planted-{os.getpid()}.lua"
    work = tmp_path / "work"
    work.mkdir()
    tex = work / "plant.tex"
    tex.write_text(r"""\documentclass{article}
\begin{document}
\directlua{
  local f = io.open("%s", "w") if f then f:write("evil") f:close() end
  local v = os.getenv("TEXMFVAR") or "" texio.write_nl("TEXMFVAR-SEEN:" .. v)
  local g = io.open(v .. "/planted-in-throwaway.lua", "w") if g then g:write("ok") g:close() end
}
x
\end{document}
""" % planted)
    try:
        r = sandbox.compile_tex("lualatex", work, "plant", tex, [work], timeout=180)
        assert r.returncode != -1, r.stderr
        log = (work / "plant.log").read_text(errors="replace").replace("\n", "")  # the log wraps long lines
        assert not planted.exists()
        assert "TEXMFVAR-SEEN:" + str(work / ".texmfvar") in log  # the engine saw the throwaway cache
        assert (work / ".texmfvar" / "planted-in-throwaway.lua").exists()  # writes there are fine and die with the run
    finally:
        if planted.exists():
            planted.unlink()


def test_symlinked_build_files_outside_the_project_are_refused(tmp_path):
    """N2: main.aux -> /somewhere/secret must not be copied into the TikZ side compile."""
    proj = tmp_path / "proj"
    proj.mkdir()
    secret = tmp_path / "secret.txt"
    secret.write_text("SECRETTOKEN42")
    src = proj / "main.tex"
    src.write_text(DOC)
    os.symlink(secret, proj / "main.aux")
    with pytest.raises(PreflightError, match="symbolic link"):
        preflight(src, need_build_files=False)
    # a symlink inside the project is fine
    (proj / "main.aux").unlink()
    (proj / "real.aux").write_text("\\relax\n")
    os.symlink(proj / "real.aux", proj / "main.aux")
    preflight(src, need_build_files=False)


def test_tikz_gets_only_the_labels_of_the_aux(tmp_path):
    from texdocxdrafter.tikz import aux_labels

    src = _write(tmp_path)
    src.with_suffix(".aux").write_text("\\relax\n\\@writefile{toc}{SECRET LINE}\n\\newlabel{fig:a}{{1}{2}}\n\\gdef\\x{SECRET}\n")
    pf = preflight(src, need_build_files=False)
    assert aux_labels(pf) == "\\newlabel{fig:a}{{1}{2}}\n"


def _template_with_private_parts(path: Path) -> None:
    """A reference document carrying a footnote, custom and extended properties, a thumbnail and core properties."""
    import docx

    d = docx.Document()
    d.add_paragraph("template text")
    cp = d.core_properties
    cp.title, cp.author, cp.subject, cp.keywords, cp.category, cp.comments = (
        "TPL-TITLE", "TPL-AUTHOR", "TPL-SUBJECT", "TPL-KEYWORDS", "TPL-CATEGORY", "TPL-COMMENT")
    d.save(str(path))
    W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    footnotes = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:footnotes xmlns:w="{W}">'
                 '<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
                 '<w:footnote w:id="1"><w:p><w:r><w:t>PRIVATE-FOOTNOTE</w:t></w:r></w:p></w:footnote></w:footnotes>')
    custom = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
              '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/custom-properties" '
              'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
              '<property fmtid="{D5CDD505-2E9C-101B-9397-08002B2CF9AE}" pid="2" name="Client"><vt:lpwstr>PRIVATE-CUSTOM</vt:lpwstr></property>'
              '</Properties>')
    tmp = path.with_suffix(".tmp")
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        has_thumb = "docProps/thumbnail.jpeg" in zin.namelist()
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "docProps/thumbnail.jpeg":
                data = b"PRIVATE-THUMB"
            if item.filename == "[Content_Types].xml":
                data = data.replace(b"</Types>", (
                    b'<Override PartName="/word/footnotes.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"/>'
                    b'<Override PartName="/docProps/custom.xml" ContentType="application/vnd.openxmlformats-officedocument.custom-properties+xml"/>'
                    b'<Default Extension="jpeg" ContentType="image/jpeg"/></Types>'))
            elif item.filename == "word/_rels/document.xml.rels":
                data = data.replace(b"</Relationships>", (
                    b'<Relationship Id="rIdFn" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes" Target="footnotes.xml"/></Relationships>'))
            elif item.filename == "_rels/.rels":
                data = data.replace(b"</Relationships>", (
                    b'<Relationship Id="rIdCu" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/custom-properties" Target="docProps/custom.xml"/>'
                    + (b"" if has_thumb else b'<Relationship Id="rIdTh" Type="http://schemas.openxmlformats.org/package/2006/relationships/metadata/thumbnail" Target="docProps/thumbnail.jpeg"/>')
                    + b'</Relationships>'))
            elif item.filename == "docProps/app.xml":
                data = data.replace(b"</Properties>", b"<Company>PRIVATE-COMPANY</Company></Properties>")
            zout.writestr(item, data)
        zout.writestr("word/footnotes.xml", footnotes)
        zout.writestr("docProps/custom.xml", custom)
        if not has_thumb:
            zout.writestr("docProps/thumbnail.jpeg", b"PRIVATE-THUMB")
    tmp.replace(path)


def test_reference_doc_private_parts_are_stripped(tmp_path):
    """N3: nothing of the template's own content or author survives; one footnotes part; valid package."""
    tpl = tmp_path / "template.docx"
    _template_with_private_parts(tpl)
    z = zipfile.ZipFile(tpl)
    assert b"PRIVATE-FOOTNOTE" in z.read("word/footnotes.xml") and "docProps/thumbnail.jpeg" in z.namelist()

    src = _write(tmp_path, DOC.replace("Hello world.", r"Hello world\footnote{paper note}."))
    res = build(src, reference_doc=tpl, no_harvest=True, log=lambda *_: None)
    assert res.check.ok, res.check.errors
    z = zipfile.ZipFile(tmp_path / "main.docx")
    names = z.namelist()
    assert names.count("word/footnotes.xml") == 1
    assert "docProps/custom.xml" not in names and "docProps/thumbnail.jpeg" not in names
    blob = b"".join(z.read(n) for n in names)
    for secret in (b"PRIVATE-FOOTNOTE", b"PRIVATE-CUSTOM", b"PRIVATE-THUMB", b"PRIVATE-COMPANY", b"TPL-TITLE",
                   b"TPL-AUTHOR", b"TPL-SUBJECT", b"TPL-KEYWORDS", b"TPL-CATEGORY", b"TPL-COMMENT"):
        assert secret not in blob, secret
    assert b"paper note" in z.read("word/footnotes.xml")


def test_symlinked_config_dir_is_refused(tmp_path):
    """N4: main.texdocxdrafter -> elsewhere must not receive any file."""
    proj = tmp_path / "proj"
    proj.mkdir()
    victim = tmp_path / "victimdir"
    victim.mkdir()
    (victim / "harvest.tex").write_text("mine")
    src = proj / "main.tex"
    src.write_text(DOC)
    src.with_suffix(".aux").write_text("\\relax\n")
    os.symlink(victim, proj / "main.texdocxdrafter")
    with pytest.raises(PreflightError, match="real directory"):
        build(src, no_harvest=True, log=lambda *_: None)
    assert sorted(p.name for p in victim.iterdir()) == ["harvest.tex"]
    assert (victim / "harvest.tex").read_text() == "mine"


CITING_DOC = r"""\documentclass{article}
\usepackage{biblatex}
\addbibresource{refs.bib}
\begin{document}
Text \cite{knuth84}.
\printbibliography
\end{document}
"""


def test_foreign_side_document_is_not_replaced_and_shipped_pdf_is_not_trusted(tmp_path, monkeypatch):
    monkeypatch.setitem(sandbox._disabled, "bwrap", True)
    src = _write(tmp_path, CITING_DOC)
    src.with_suffix(".bbl").write_text("% fake\n")
    (tmp_path / "refs.bib").write_text(BIB)
    cfg = tmp_path / "main.texdocxdrafter"
    cfg.mkdir()
    (cfg / "harvest.tex").write_text("someone else's file")
    with pytest.raises(PreflightError, match="not written by texdocxdrafter"):
        build(src, log=lambda *_: None)
    assert (cfg / "harvest.tex").read_text() == "someone else's file"
    (cfg / "harvest.tex").unlink()
    # a harvest.tex with exactly our content plus a PDF, shipped in an archive without our registry
    build(src, log=lambda *_: None)
    ours = (cfg / "harvest.tex").read_text()
    (cfg / "outputs.json").unlink()
    time.sleep(0.05)
    (cfg / "harvest.pdf").write_bytes(b"%PDF-1.4 not really")
    res = build(src, overwrite=True, log=lambda *_: None)  # overwrite: without the registry our own docx counts as foreign
    msgs = "\n".join(e.message for e in res.audit.entries)
    assert "for you to compile" in msgs  # the shipped PDF was not used
    assert (cfg / "harvest.tex").read_text() == ours
    assert (cfg / "harvest.tex").stat().st_mtime >= (cfg / "harvest.pdf").stat().st_mtime


def test_compile_command_quotes_paths(tmp_path):
    """N5: a project directory named to inject shell commands is printed harmlessly."""
    import shlex
    import subprocess

    from texdocxdrafter.sidecompile import compile_command

    marker = tmp_path / "INJECTED"
    proj = tmp_path / "proj$(touch INJECTED) it's"
    proj.mkdir()
    src = proj / "main.tex"
    src.write_text(DOC)
    pf = preflight(src, need_build_files=False)
    cmd = compile_command(pf, proj / "main.texdocxdrafter", "harvest", "biber")
    tokens = shlex.split(cmd)
    assert tokens[:2] == ["cd", str(proj / "main.texdocxdrafter")]
    assert any(t.startswith("TEXINPUTS=") and t.endswith(f"{proj}:") for t in tokens)
    subprocess.run(cmd.split(" && ")[0], shell=True, cwd=tmp_path, capture_output=True, timeout=30)  # just the cd part
    assert not marker.exists()


def test_inventory_does_not_write_through_a_dangling_symlink(tmp_path, capsys):
    """N6: stylemap.yaml -> missing target must not create that target."""
    from texdocxdrafter import cli
    from texdocxdrafter.stylemap import dump_skeleton

    src = _write(tmp_path)
    cfg = tmp_path / "main.texdocxdrafter"
    cfg.mkdir()
    victim = tmp_path / "victim-created-by-inventory"
    os.symlink(victim, cfg / "stylemap.yaml")
    assert cli.main(["inventory", str(src)]) == 2
    assert not victim.exists()
    with pytest.raises(FileExistsError):
        dump_skeleton({}, cfg / "stylemap.yaml")
    assert not victim.exists()


def test_only_web_and_mail_links_become_hyperlinks(tmp_path):
    src = _write(tmp_path, DOC.replace("Hello world.", r"See \url{https://ok.example/x}, \href{mailto:a@b.c}{mail}, "
                                                        r"\url{file:///etc/passwd} and \href{javascript:alert(1)}{click}."))
    build(src, no_harvest=True, log=lambda *_: None)
    z = zipfile.ZipFile(tmp_path / "main.docx")
    rels = z.read("word/_rels/document.xml.rels").decode()
    assert "https://ok.example/x" in rels and "mailto:a@b.c" in rels
    assert "file:" not in rels and "javascript:" not in rels
    xml = z.read("word/document.xml").decode()
    assert "file:///etc/passwd" in xml and "click" in xml  # kept as visible text


# ---- fifth review round --------------------------------------------------------------------------

LUA_RELATIVE_SYMLINK = r"""\documentclass{article}
\begin{document}
\directlua{
  lfs.mkdir("a") lfs.mkdir("a/b") lfs.mkdir("a/b/c")
  local r0, e0 = lfs.link("../victim.txt", "harvest.txt", true)
  lfs.chdir("a/b/c")
  % checked from the current directory the target looks like work/victim.txt; the kernel resolves it from the
  % link's own directory, where it is ../victim.txt next to work
  local r1, e1 = lfs.link("../victim.txt", "../../../harvest.txt", true)
  local r2, e2 = lfs.chdir("/")
  texio.write_nl("LINKS:" .. tostring(r0) .. ":" .. tostring(r1) .. ":" .. tostring(r2))
}
x
\end{document}
"""


@pytest.mark.skipif(not HAVE_LUALATEX, reason="lualatex not installed")
def test_lua_cannot_plant_a_symbolic_link(tmp_path, monkeypatch):
    """N7: a relative symlink target evaluated after chdir slipped past the guard; no symlinks at all now."""
    monkeypatch.setitem(sandbox._disabled, "bwrap", True)
    (tmp_path / "victim.txt").write_text("untouched")
    work = tmp_path / "work"
    work.mkdir()
    tex = work / "evil.tex"
    tex.write_text(LUA_RELATIVE_SYMLINK)
    r = sandbox.compile_tex("lualatex", work, "evil", tex, [work], timeout=120, check_policy=False)
    assert r.returncode != -1, r.stderr
    log = (work / "evil.log").read_text(errors="replace").replace("\n", "")
    assert "LINKS:nil:nil:nil" in log, log[-800:]
    assert not (work / "harvest.txt").is_symlink() and not (work / "harvest.txt").exists()
    assert not (work / "a" / "b" / "c" / "harvest.txt").exists()


def test_work_dir_helpers_never_follow_a_symlink(tmp_path):
    victim = tmp_path / "victim.txt"
    victim.write_text("untouched")
    secret = tmp_path / "secret.txt"
    secret.write_text("TOPSECRET")
    work = tmp_path / "work"
    work.mkdir()
    (work / "out.txt").symlink_to(victim)
    (work / "log.txt").symlink_to(secret)
    (work / "png").symlink_to(tmp_path)
    assert sandbox.read_untrusted(work / "log.txt") is None and sandbox.read_untrusted_text(work / "log.txt") == ""
    assert not sandbox.is_plain_file(work / "log.txt") and sandbox.is_plain_file(secret)
    sandbox.write_fresh(work / "out.txt", "ours")
    assert victim.read_text() == "untouched"
    assert sandbox.is_plain_file(work / "out.txt") and (work / "out.txt").read_text() == "ours"
    d = sandbox.fresh_dir(work / "png")
    assert d.is_dir() and not d.is_symlink() and sorted(p.name for p in d.iterdir()) == []
    assert secret.exists() and victim.exists()  # the symlinked directory's content was not removed


def _fake_compile(plant):
    """A compile_tex stand-in: the 'document' leaves whatever ``plant(work)`` creates, and reports success."""
    import subprocess

    def compile_tex(engine, work, jobname, texfile, dirs, timeout=0, check_policy=True):
        plant(Path(work))
        return subprocess.CompletedProcess([engine], 0, "", "")

    return compile_tex


def test_planted_symlinks_in_the_harvest_work_dir_are_not_followed(tmp_path, monkeypatch):
    """N7: the tool writes harvest.txt and reads harvest.log after the compile; neither goes through a link."""
    victim = tmp_path / "victim.txt"
    victim.write_text("untouched")
    secret = tmp_path / "secret.txt"
    secret.write_text("TOPSECRET")
    src = _write(tmp_path, CITING_DOC)
    src.with_suffix(".bbl").write_text("% fake\n")
    (tmp_path / "refs.bib").write_text(BIB)

    def plant(work):  # called once per engine run
        if not (work / "harvest.log").is_symlink():
            (work / "harvest.log").symlink_to(secret)
            (work / "harvest.pdf").write_bytes(b"%PDF-1.4 fake")
            (work / "harvest.txt").symlink_to(victim)

    real_run = sandbox.run

    def run(cmd, *a, **k):
        if cmd[0] == "pdftotext":
            import subprocess
            return subprocess.CompletedProcess(cmd, 0, "@@CITE 1@@harvested@@END@@\n", "")
        return real_run(cmd, *a, **k)

    monkeypatch.setattr(sandbox, "have_os_sandbox", lambda: True)
    monkeypatch.setattr(sandbox, "compile_tex", _fake_compile(plant))
    monkeypatch.setattr(sandbox, "run", run)
    res = build(src, keep_work=True, log=lambda *_: None)
    assert victim.read_text() == "untouched"
    report = (tmp_path / "main.texdocxdrafter-report.md").read_text()
    assert "TOPSECRET" not in report
    txt = tmp_path / "main.texdocxdrafter" / "work" / "harvest.txt"
    assert txt.is_file() and not txt.is_symlink() and "harvested" in txt.read_text()


TIKZ_DOC = r"""\documentclass{article}
\usepackage{tikz}
\begin{document}
A picture: \begin{tikzpicture}\draw (0,0) -- (1,1);\end{tikzpicture}
\end{document}
"""


def test_planted_tikz_outputs_are_not_followed(tmp_path, monkeypatch):
    """N7: a symlinked tikz.pdf, tikz.log or png directory left by the compile is ignored, not embedded."""
    secret = tmp_path / "secret.txt"
    secret.write_text("TOPSECRET")
    secret_png = tmp_path / "pngs"
    secret_png.mkdir()
    (secret_png / "page-1.png").write_bytes(b"\x89PNG not really")
    src = _write(tmp_path, TIKZ_DOC)

    def plant(work):
        (work / "tikz.log").symlink_to(secret)
        (work / "tikz.pdf").symlink_to(secret)
        (work / "png").symlink_to(secret_png)

    monkeypatch.setattr(sandbox, "have_os_sandbox", lambda: True)
    monkeypatch.setattr(sandbox, "compile_tex", _fake_compile(plant))
    res = build(src, no_harvest=True, log=lambda *_: None)
    msgs = "\n".join(e.message for e in res.audit.entries)
    assert "tikz compile failed" in msgs and "TOPSECRET" not in msgs
    assert not [n for n in zipfile.ZipFile(tmp_path / "main.docx").namelist() if n.startswith("word/media/")]
    assert (secret_png / "page-1.png").exists()


@pytest.mark.skipif(not HAVE_PDFLATEX, reason="pdflatex not installed")
def test_sandbox_shows_neither_etc_nor_proc(tmp_path):
    """N9: only the few /etc entries TeX needs are visible; a document cannot typeset /etc/passwd."""
    if not sandbox.have_os_sandbox():
        pytest.skip("bubblewrap not available")
    assert sandbox.run(["cat", "/etc/passwd"], cwd=tmp_path, rw=[], ro=[tmp_path]).returncode != 0
    assert sandbox.run(["ls", "/proc"], cwd=tmp_path, rw=[], ro=[tmp_path]).returncode != 0
    if Path("/etc/fonts").is_dir():
        assert sandbox.run(["ls", "/etc/fonts"], cwd=tmp_path, rw=[], ro=[tmp_path]).returncode == 0
    work = tmp_path / "work"
    work.mkdir()
    tex = work / "peek.tex"
    tex.write_text(r"""\documentclass{article}
\begin{document}
\InputIfFileExists{/etc/passwd}{\typeout{READ-OK}}{\typeout{READ-REFUSED}}
\InputIfFileExists{/etc/hostname}{\typeout{READ-OK}}{\typeout{READ-REFUSED}}
x
\end{document}
""")
    sandbox.compile_tex("pdflatex", work, "peek", tex, [work], timeout=120)
    log = (work / "peek.log").read_text(errors="replace")
    assert log.count("READ-REFUSED") == 2 and "READ-OK" not in log


def _template_with_header_fields(path: Path) -> None:
    import docx
    from docx.oxml import parse_xml

    W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
    d = docx.Document()
    p = d.sections[0].header.paragraphs[0]
    p.text = "Running head "
    for xml in (
        f'<w:fldSimple {W} w:instr=" PAGE "><w:r><w:t>1</w:t></w:r></w:fldSimple>',
        f'<w:fldSimple {W} w:instr=" INCLUDEPICTURE &quot;http://evil.example/x.png&quot; \\d "><w:r><w:t>pic</w:t></w:r></w:fldSimple>',
        f'<w:r {W}><w:fldChar w:fldCharType="begin"/></w:r>',
        f'<w:r {W}><w:instrText xml:space="preserve"> INCLUDETEXT "\\\\\\\\evil\\\\share\\\\notes.docx" </w:instrText></w:r>',
        f'<w:r {W}><w:fldChar w:fldCharType="separate"/></w:r>',
        f'<w:r {W}><w:t>INCLUDED-SECRET</w:t></w:r>',
        f'<w:r {W}><w:fldChar w:fldCharType="end"/></w:r>',
        f'<w:r {W}><w:fldChar w:fldCharType="begin"/></w:r>',
        f'<w:r {W}><w:instrText xml:space="preserve"> NUMPAGES </w:instrText></w:r>',
        f'<w:r {W}><w:fldChar w:fldCharType="separate"/></w:r>',
        f'<w:r {W}><w:t>9</w:t></w:r>',
        f'<w:r {W}><w:fldChar w:fldCharType="end"/></w:r>',
    ):
        p._p.append(parse_xml(xml))
    d.save(str(path))


def test_reference_doc_header_fields_that_fetch_content_are_removed(tmp_path):
    """N8: INCLUDETEXT/INCLUDEPICTURE in the template's header go; PAGE and NUMPAGES stay."""
    tpl = tmp_path / "template.docx"
    _template_with_header_fields(tpl)
    src = _write(tmp_path)
    res = build(src, reference_doc=tpl, no_harvest=True, log=lambda *_: None)
    assert res.check.ok, res.check.errors
    z = zipfile.ZipFile(tmp_path / "main.docx")
    headers = "".join(z.read(n).decode("utf-8") for n in z.namelist() if n.startswith("word/header"))
    assert "Running head" in headers and " PAGE " in headers and "NUMPAGES" in headers and ">9<" in headers
    for gone in ("INCLUDETEXT", "INCLUDEPICTURE", "evil", "INCLUDED-SECRET", "notes.docx"):
        assert gone not in headers, gone
    assert "w:fldChar" in headers and headers.count('w:fldCharType="begin"') == headers.count('w:fldCharType="end"') == 1


def _bounded(fn, seconds=90):
    import signal

    def alarm(*_):
        raise AssertionError(f"did not finish within {seconds} s")

    old = signal.signal(signal.SIGALRM, alarm)
    signal.alarm(seconds)
    try:
        return fn()
    finally:
        signal.alarm(0)
        signal.signal(signal.SIGALRM, old)


RUNAWAY = {
    "macro doubling": r"\newcommand{\dbl}{\dbl\dbl}" + "\n" + r"\begin{document}x \dbl y\end{document}",
    "environment doubling": r"\newenvironment{dd}{\begin{dd}\begin{dd}}{\end{dd}\end{dd}}" + "\n"
                            + r"\begin{document}\begin{dd}x\end{dd}\end{document}",
    "math argument doubling": r"\newcommand{\twice}[1]{#1#1}" + "\n" + r"\begin{document}$"
                              + r"\twice{" * 30 + "x" + "}" * 30 + r"$\end{document}",
}


@pytest.mark.parametrize("name", sorted(RUNAWAY))
def test_runaway_macro_expansion_is_bounded(tmp_path, name):
    """N10: exponential expansion stops at the budget, with an error in the report, within seconds."""
    src = _write(tmp_path, "\\documentclass{article}\n" + RUNAWAY[name] + "\n")
    res = _bounded(lambda: build(src, no_harvest=True, no_tikz=True, log=lambda *_: None))
    assert res.check.ok, res.check.errors
    errors = [e.message for e in res.audit.entries if e.severity == "error"]
    assert errors and all("expansion stopped" in m or "depth limit" in m for m in errors), errors
    assert len(res.audit.warnings()) < 20, len(res.audit.warnings())  # one note, not one item per leftover macro


def test_deep_nesting_is_an_error_not_a_traceback(tmp_path, capsys):
    """N10: 3000 nested braces exceed Python's recursion depth; the CLI reports it and writes nothing."""
    from texdocxdrafter.cli import main

    src = _write(tmp_path, "\\documentclass{article}\n\\begin{document}\n" + "{" * 3000 + "x" + "}" * 3000
                 + "\n\\end{document}\n")
    rc = main(["build", str(src), "--no-harvest", "--no-tikz"])
    assert rc == 1
    assert "nests groups" in capsys.readouterr().err
    assert not (tmp_path / "main.docx").exists()


def test_tool_process_has_memory_and_cpu_limits():
    import subprocess
    import sys

    code = ("import resource; from texdocxdrafter import sandbox; sandbox.limit_self(); "
            "print(resource.getrlimit(resource.RLIMIT_AS)[0], resource.getrlimit(resource.RLIMIT_CPU)[0])")
    out = subprocess.run([sys.executable, "-c", code], capture_output=True, text=True, check=True).stdout.split()
    assert int(out[0]) <= sandbox.TOOL_MAX_MEMORY and int(out[1]) <= sandbox.TOOL_CPU_SECONDS


# ---- sixth review round --------------------------------------------------------------------------

def test_sandbox_drops_capabilities_and_limits_children_without_preexec(tmp_path):
    import sys

    if sandbox.have_os_sandbox():
        args = sandbox._wrap(["true"], [tmp_path], [])
        assert args[args.index("--cap-drop") + 1] == "ALL"
    if not sandbox._prlimit():
        pytest.skip("prlimit not available")
    code = "import resource; print(resource.getrlimit(resource.RLIMIT_AS)[0], resource.getrlimit(resource.RLIMIT_CPU)[0])"
    r = sandbox.run([sys.executable, "-c", code], cwd=tmp_path, timeout=30, max_memory=1 << 30)  # no bwrap: rw not given
    assert r.stdout.split() == [str(1 << 30), "30"], r.stderr


def _rezip(path: Path, replace: dict, add: dict) -> None:
    tmp = path.with_suffix(".tmp")
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            for old, new in replace.get(item.filename, ()):
                assert old in data, (item.filename, old)
                data = data.replace(old, new)
            zout.writestr(item, data)
        for name, data in add.items():
            zout.writestr(name, data)
    tmp.replace(path)


VBA_REL = b'<Relationship Id="rIdVba" Type="http://schemas.microsoft.com/office/2006/relationships/vbaProject" Target="vbaProject.bin"/>'
VBA_CT = b'<Override PartName="/word/vbaProject.bin" ContentType="application/vnd.ms-office.vbaProject"/>'


def test_macro_enabled_template_is_refused(tmp_path):
    """N8 residual: a .docm/.dotm template is refused outright rather than silently carrying its VBA."""
    import docx

    from texdocxdrafter.refdoc import TemplateError

    tpl = tmp_path / "template.docm"
    docx.Document().save(str(tpl))
    _rezip(tpl, {
        "[Content_Types].xml": [(b"application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml",
                                 b"application/vnd.ms-word.document.macroEnabled.main+xml"), (b"</Types>", VBA_CT + b"</Types>")],
        "word/_rels/document.xml.rels": [(b"</Relationships>", VBA_REL + b"</Relationships>")],
    }, {"word/vbaProject.bin": b"VBA-SECRET"})
    src = _write(tmp_path)
    with pytest.raises(TemplateError, match="macro-enabled"):
        build(src, reference_doc=tpl, no_harvest=True, log=lambda *_: None)
    assert not (tmp_path / "main.docx").exists()


def test_vba_and_embedded_objects_in_template_are_stripped(tmp_path):
    """N8 residual: a VBA project smuggled into a .docx and an OLE object in the header do not travel."""
    import docx
    from docx.oxml import parse_xml

    tpl = tmp_path / "template.docx"
    d = docx.Document()
    p = d.sections[0].header.paragraphs[0]
    p.text = "Running head "
    W = ('xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
         'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
         'xmlns:o="urn:schemas-microsoft-com:office:office" xmlns:v="urn:schemas-microsoft-com:vml"')
    p._p.append(parse_xml(f'<w:r {W}><w:object><v:shape id="s1" style="width:10pt;height:10pt"/>'
                          '<o:OLEObject Type="Embed" ProgID="Package" ShapeID="s1" DrawAspect="Icon" ObjectID="_1" r:id="rIdOle"/>'
                          '</w:object></w:r>'))
    d.save(str(tpl))
    with zipfile.ZipFile(tpl) as z:
        header = [n for n in z.namelist() if n.startswith("word/header") and n.endswith(".xml")][0]
        header_rels = "word/_rels/" + header.split("/")[-1] + ".rels"
        has_rels = header_rels in z.namelist()
    ole_rel = (b'<Relationship Id="rIdOle" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/oleObject" '
               b'Target="embeddings/oleObject1.bin"/>')
    add = {"word/vbaProject.bin": b"VBA-SECRET", "word/embeddings/oleObject1.bin": b"OLE-SECRET"}
    replace = {
        "[Content_Types].xml": [(b"</Types>", VBA_CT + b'<Default Extension="bin" ContentType="application/vnd.openxmlformats-officedocument.oleObject"/></Types>')],
        "word/_rels/document.xml.rels": [(b"</Relationships>", VBA_REL + b"</Relationships>")],
    }
    if has_rels:
        replace[header_rels] = [(b"</Relationships>", ole_rel + b"</Relationships>")]
    else:
        add[header_rels] = (b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships '
                            b'xmlns="http://schemas.openxmlformats.org/package/2006/relationships">' + ole_rel + b"</Relationships>")
    _rezip(tpl, replace, add)
    src = _write(tmp_path)
    res = build(src, reference_doc=tpl, no_harvest=True, log=lambda *_: None)
    assert res.check.ok, res.check.errors
    z = zipfile.ZipFile(tmp_path / "main.docx")
    blob = b"".join(z.read(n) for n in z.namelist())
    assert b"VBA-SECRET" not in blob and b"OLE-SECRET" not in blob and b"vbaProject" not in blob
    headers = "".join(z.read(n).decode() for n in z.namelist() if n.startswith("word/header"))
    assert "Running head" in headers and "w:object" not in headers and "OLEObject" not in headers


def test_style_names_are_escaped_and_validated(tmp_path):
    """N11: a style name with quotes, ampersand and angle brackets is escaped; control characters are refused."""
    import docx

    from texdocxdrafter.stylemap import StyleMapError, validate

    name = 'Body "Text" & <Co>'
    sm = tmp_path / "map.yaml"
    sm.write_text(f"paragraph:\n  body: '{name}'\n")
    src = _write(tmp_path)
    res = build(src, style_map=sm, no_harvest=True, log=lambda *_: None)
    assert res.check.ok, res.check.errors
    from docx.oxml.ns import qn

    styles = docx.Document(str(tmp_path / "main.docx")).styles.element  # styles[name] itself cannot quote such a name
    assert any(st.find(qn("w:name")).get(qn("w:val")) == name for st in styles.findall(qn("w:style")))
    with pytest.raises(StyleMapError, match="printable"):
        validate({"paragraph": {"body": "Body\x00Text"}})
    with pytest.raises(StyleMapError, match="printable"):
        validate({"character": {"emph": "x" * 300}})


def test_libreoffice_validation_runs_sandboxed(tmp_path, monkeypatch):
    """N12: the validation run gets the same sandbox as every other tool, with a private profile."""
    import subprocess

    from texdocxdrafter import pipeline as pl

    seen = {}

    def run(cmd, cwd, env=None, timeout=0, rw=None, ro=None, **k):
        seen.update(cmd=cmd, env=env, rw=rw, ro=ro)
        return subprocess.CompletedProcess(cmd, 1, "", "no office here")

    monkeypatch.setattr(pl.shutil, "which", lambda name: "/opt/fake/soffice" if name == "soffice" else None)
    monkeypatch.setattr(sandbox, "run", run)
    src = _write(tmp_path)
    res = build(src, no_harvest=True, validate_lo=True, log=lambda *_: None)
    # LibreOffice reads the private candidate in the work directory, never the output directory itself
    assert seen["rw"] and seen["env"]["HOME"] == str(seen["rw"][0])
    assert len(seen["ro"]) == 1 and seen["ro"][0].name.startswith("texdocxdrafter-") and seen["ro"][0] != tmp_path
    assert any(a.startswith("-env:UserInstallation=file://") for a in seen["cmd"])
    assert any("inside the sandbox" in e.message for e in res.audit.entries)


def test_bookmark_names_use_sha256():
    import hashlib

    from texdocxdrafter.writer import bm_name

    key = "a-very-long-citation-key-" * 4
    name = bm_name("cite_", key)
    assert len(name) == 40 and name.endswith(hashlib.sha256(key.encode()).hexdigest()[:9])


def test_build_limits_its_own_process_and_reports_when_it_cannot(tmp_path, monkeypatch):
    """N10 residual: build() applies the limits itself, so a library caller gets them too."""
    calls = []
    monkeypatch.setattr(sandbox, "limit_self", lambda *a, **k: calls.append(1) or ["memory 4 GiB (test)"])
    src = _write(tmp_path)
    res = build(src, no_harvest=True, log=lambda *_: None)
    assert calls and any("could not limit the tool's own process: memory 4 GiB (test)" in e.message for e in res.audit.entries)
    assert sandbox.limit_self.__name__ == "<lambda>"  # the real one is exercised by test_tool_process_has_memory_and_cpu_limits


# ---- seventh review round ------------------------------------------------------------------------

def test_mail_merge_source_and_mergefields_do_not_survive_the_template(tmp_path):
    """N14: MERGEFIELD in the header plus a w:mailMerge data source would make Word connect on update."""
    import docx
    from docx.oxml import parse_xml

    W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
    tpl = tmp_path / "template.docx"
    d = docx.Document()
    p = d.sections[0].header.paragraphs[0]
    p.text = "Dear "
    for xml in (f'<w:r {W}><w:fldChar w:fldCharType="begin"/></w:r>',
                f'<w:r {W}><w:instrText xml:space="preserve"> MERGEFIELD Name </w:instrText></w:r>',
                f'<w:r {W}><w:fldChar w:fldCharType="separate"/></w:r>', f'<w:r {W}><w:t>«Name»</w:t></w:r>',
                f'<w:r {W}><w:fldChar w:fldCharType="end"/></w:r>'):
        p._p.append(parse_xml(xml))
    d.save(str(tpl))
    mm = (b'<w:mailMerge><w:mainDocumentType w:val="formLetters"/><w:dataType w:val="database"/>'
          b'<w:dataSource r:id="rIdMM"/></w:mailMerge>')
    rel = (b'<Relationship Id="rIdMM" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/mailMergeSource" '
           b'Target="file:///\\\\evil.example\\share\\data.mdb" TargetMode="External"/>')
    with zipfile.ZipFile(tpl) as z:
        settings = z.read("word/settings.xml")
        has_rels = "word/_rels/settings.xml.rels" in z.namelist()
    r_ns = b'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
    first_child = settings.index(b">", settings.index(b"<w:settings")) + 1
    patched = settings[:first_child] + mm + settings[first_child:]
    if r_ns not in patched:
        patched = patched.replace(b"<w:settings ", b"<w:settings " + r_ns + b" ", 1)
    replace = {"word/settings.xml": [(settings, patched)]}
    add = {}
    if has_rels:
        replace["word/_rels/settings.xml.rels"] = [(b"</Relationships>", rel + b"</Relationships>")]
    else:
        add["word/_rels/settings.xml.rels"] = (b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships '
                                               b'xmlns="http://schemas.openxmlformats.org/package/2006/relationships">' + rel + b"</Relationships>")
    _rezip(tpl, replace, add)
    src = _write(tmp_path)
    res = build(src, reference_doc=tpl, no_harvest=True, log=lambda *_: None)
    assert res.check.ok, res.check.errors
    z = zipfile.ZipFile(tmp_path / "main.docx")
    blob = b"".join(z.read(n) for n in z.namelist())
    assert b"mailMerge" not in blob and b"evil.example" not in blob and b"MERGEFIELD" not in blob
    headers = "".join(z.read(n).decode() for n in z.namelist() if n.startswith("word/header"))
    assert "Dear" in headers


def test_cross_including_files_do_not_fan_out(tmp_path):
    """N18: 12 files that each \\input all the others are O(N!) chains; the inclusion budget stops it."""
    names = [f"part{i}" for i in range(12)]
    for n in names:
        (tmp_path / f"{n}.tex").write_text("".join(f"\\input{{{m}}}\n" for m in names) + f"text of {n}\n")
    src = _write(tmp_path, "\\documentclass{article}\n\\begin{document}\n" + "".join(f"\\input{{{m}}}\n" for m in names)
                 + "\\end{document}\n")
    pf = _bounded(lambda: preflight(src, need_build_files=False), 60)
    assert any("million characters included" in i or "\\input files" in i for i in pf.issues), pf.issues[:3]


def test_column_multiplier_is_bounded():
    from texdocxdrafter.parser.tables import MAX_COLUMNS, parse_colspec

    cols = _bounded(lambda: parse_colspec("*{999999999}{l}"), 30)
    assert len(cols) == MAX_COLUMNS
    assert len(parse_colspec("*{3}{lc}|r")) == 7


def test_keep_work_never_follows_a_symlink_at_work(tmp_path):
    """N20: main.texdocxdrafter/work -> elsewhere: the target is neither deleted nor written into."""
    victim = tmp_path / "victim"
    victim.mkdir()
    (victim / WORK_MARKER).write_text("looks like ours")
    (victim / "precious.txt").write_text("keep me")
    src = _write(tmp_path)
    cfg = tmp_path / "main.texdocxdrafter"
    cfg.mkdir()
    (cfg / "work").symlink_to(victim)
    build(src, no_harvest=True, keep_work=True, log=lambda *_: None)
    assert (victim / "precious.txt").read_text() == "keep me" and (cfg / "work").is_symlink()
    assert [p for p in cfg.iterdir() if p.name.startswith("work-")]


def test_tree_helpers_do_not_recurse_without_bound():
    """N21: walk_tree is iterative; a bbl optional argument with thousands of nested braces is read."""
    from texdocxdrafter.bbl import _balanced
    from texdocxdrafter.parser.model import Text, walk_tree

    deep = Text(text="leaf")
    for _ in range(5000):
        deep = [deep]
    seen = []
    walk_tree(deep, seen.append)
    assert len(seen) == 1
    inner, end = _balanced("[" + "{" * 5000 + "x" + "}" * 5000 + "]tail", 0, "[", "]")
    assert inner.endswith("x" + "}" * 5000) and end == 10003


def test_thousands_of_declarations_in_a_row_convert(tmp_path):
    """N22: consecutive \\bfseries no longer copies the tail each time nor nests without bound."""
    src = _write(tmp_path, "\\documentclass{article}\n\\begin{document}\n" + "\\bfseries " * 5000 + "bold at last\n\\end{document}\n")
    res = _bounded(lambda: build(src, no_harvest=True, no_tikz=True, log=lambda *_: None), 60)
    assert res.check.ok, res.check.errors
    assert "bold at last" in _docx_xml(tmp_path / "main.docx")


def test_cli_reports_unmet_process_limits_for_every_command(tmp_path, monkeypatch, capsys):
    from texdocxdrafter.cli import main

    monkeypatch.setattr(sandbox, "limit_self", lambda *a, **k: ["memory 4 GiB (test)"])
    src = _write(tmp_path)
    assert main(["parse", str(src), "--facts"]) == 0
    assert "could not limit the tool's own process: memory 4 GiB (test)" in capsys.readouterr().err


def test_timeout_option_has_an_upper_bound(tmp_path, capsys):
    from texdocxdrafter.cli import MAX_TIMEOUT, main

    src = _write(tmp_path)
    with pytest.raises(SystemExit) as exc:
        main(["build", str(src), "--timeout", str(MAX_TIMEOUT + 1)])
    assert exc.value.code == 2 and "between 1 and" in capsys.readouterr().err


# ---- eighth review round -------------------------------------------------------------------------

def _settings_patch(tpl: Path, inject: bytes, rel: bytes | None = None) -> None:
    """Insert ``inject`` as the first children of w:settings (and a settings relationship)."""
    with zipfile.ZipFile(tpl) as z:
        settings = z.read("word/settings.xml")
        has_rels = "word/_rels/settings.xml.rels" in z.namelist()
    r_ns = b'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
    first_child = settings.index(b">", settings.index(b"<w:settings")) + 1
    patched = settings[:first_child] + inject + settings[first_child:]
    if r_ns not in patched:
        patched = patched.replace(b"<w:settings ", b"<w:settings " + r_ns + b" ", 1)
    replace = {"word/settings.xml": [(settings, patched)]}
    add = {}
    if rel:
        if has_rels:
            replace["word/_rels/settings.xml.rels"] = [(b"</Relationships>", rel + b"</Relationships>")]
        else:
            add["word/_rels/settings.xml.rels"] = (b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships '
                                                   b'xmlns="http://schemas.openxmlformats.org/package/2006/relationships">' + rel + b"</Relationships>")
    _rezip(tpl, replace, add)


def test_template_background_fill_and_protection_do_not_survive(tmp_path):
    """N25, N29: w:background (beside the body) with an external fill, displayBackgroundShape, protection."""
    import docx

    tpl = tmp_path / "template.docx"
    docx.Document().save(str(tpl))
    with zipfile.ZipFile(tpl) as z:
        document = z.read("word/document.xml")
    bg = (b'<w:background w:color="FFFFFF"><v:background xmlns:v="urn:schemas-microsoft-com:vml">'
          b'<v:fill src="http://evil.example/track.png" type="frame"/></v:background></w:background>')
    body_at = document.index(b"<w:body")
    _rezip(tpl, {"word/document.xml": [(document, document[:body_at] + bg + document[body_at:])]}, {})
    _settings_patch(tpl, b'<w:writeProtection w:recommended="1"/><w:view w:val="web"/>'
                         b'<w:documentProtection w:edit="readOnly" w:enforcement="1"/><w:displayBackgroundShape/>')
    src = _write(tmp_path)
    res = build(src, reference_doc=tpl, no_harvest=True, log=lambda *_: None)
    assert res.check.ok, res.check.errors
    z = zipfile.ZipFile(tmp_path / "main.docx")
    blob = b"".join(z.read(n) for n in z.namelist())
    for gone in (b"w:background", b"evil.example", b"displayBackgroundShape", b"documentProtection", b"writeProtection",
                 b'w:view w:val="web"'):
        assert gone not in blob, gone


def test_template_web_frameset_is_removed(tmp_path):
    """N28: a frameset in webSettings.xml names source files Word would load in web layout."""
    import docx

    tpl = tmp_path / "template.docx"
    docx.Document().save(str(tpl))
    W = b'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
    with zipfile.ZipFile(tpl) as z:
        names = z.namelist()
    if "word/webSettings.xml" in names:
        with zipfile.ZipFile(tpl) as z:
            web = z.read("word/webSettings.xml")
        first_child = web.index(b">", web.index(b"<w:webSettings")) + 1
        _rezip(tpl, {"word/webSettings.xml": [(web, web[:first_child] + b'<w:frameset><w:frame><w:sourceFileName r:id="rIdF"/></w:frame></w:frameset>' + web[first_child:])]}, {})
    else:
        _rezip(tpl, {
            "[Content_Types].xml": [(b"</Types>", b'<Override PartName="/word/webSettings.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.webSettings+xml"/></Types>')],
            "word/_rels/document.xml.rels": [(b"</Relationships>", b'<Relationship Id="rIdWeb" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/webSettings" Target="webSettings.xml"/></Relationships>')],
        }, {"word/webSettings.xml": b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:webSettings ' + W
            + b' xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><w:frameset><w:frame>'
            b'<w:sourceFileName r:id="rIdF"/></w:frame></w:frameset></w:webSettings>',
            "word/_rels/webSettings.xml.rels": b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships '
            b'xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rIdF" '
            b'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/frame" Target="http://evil.example/frame.htm" TargetMode="External"/></Relationships>'})
    src = _write(tmp_path)
    res = build(src, reference_doc=tpl, no_harvest=True, log=lambda *_: None)
    assert res.check.ok, res.check.errors
    z = zipfile.ZipFile(tmp_path / "main.docx")
    blob = b"".join(z.read(n) for n in z.namelist())
    assert b"frameset" not in blob and b"evil.example" not in blob


def test_bib_text_cannot_forge_harvest_markers(tmp_path):
    """N26: markers carry a per-run nonce; '@@END@@' or '@@ENTRY x@@' inside an entry is just text."""
    from texdocxdrafter.audit import Audit
    from texdocxdrafter.harvest import parse_text

    n = "0badf00d"
    text = (f"@@{n}CITE 0@@{n} Smith @@END@@ not the end (2020)@@{n}END@@{n}\n"
            f"@@{n}BIB 0 0@@{n}\n@@{n}TITLE@@{n} References @@{n}/TITLE@@{n}\n"
            f"@@{n}ENTRY smith20@@{n} Smith, J. (2020). @@ENTRY injected@@ Fake entry. @@{n}I@@{n}Title@@{n}/I@@{n}.\n"
            f"@@{n}ENDBIB@@{n}\n")
    res = parse_text(text, 1, [(0, 0)], Audit(), nonce=n)
    assert "".join(r.text for r in res.cites[0]) == "Smith @@END@@ not the end (2020)"
    keys = [k for k, _ in res.bibs[(0, 0)]]
    assert keys == ["smith20"]
    assert "@@ENTRY injected@@ Fake entry" in "".join(r.text for r in res.bibs[(0, 0)][0][1])
    # without the nonce the same text would have been cut and extended
    res0 = parse_text(text.replace(n, ""), 1, [(0, 0)], Audit())
    assert [k for k, _ in res0.bibs[(0, 0)]] == ["smith20", "injected"]


def test_harvest_document_uses_a_nonce_and_keeps_it_for_user_compiles(tmp_path, monkeypatch):
    from texdocxdrafter.harvest import build_harvest_tex
    from texdocxdrafter.pipeline import Registry
    from texdocxdrafter.sidecompile import SideDir, harvest_salt

    src = _write(tmp_path, CITING_DOC)
    src.with_suffix(".bbl").write_text("% fake\n")
    (tmp_path / "refs.bib").write_text(BIB)
    pf = preflight(src, need_build_files=True)
    from texdocxdrafter.audit import Audit
    from texdocxdrafter.parser import parse_document

    doc = parse_document(pf, Audit(), pipeline.stylemap_mod.load(None))
    tex, _ = build_harvest_tex(pf, doc, "0badf00d")
    assert "@@0badf00dCITE 0@@0badf00d" in tex and "@@END@@" not in tex
    cfg = tmp_path / "main.texdocxdrafter"
    cfg.mkdir()
    side = SideDir(cfg, Registry(cfg, pf))
    a = harvest_salt(pf, side)
    side.registry.save()
    b = harvest_salt(pf, SideDir(cfg, Registry(cfg, pf)))
    assert a == b and len(a) == 8
    (cfg / "harvest.salt").write_text("deadbeef\n")  # a valid salt shipped without the registry is used as is
    assert harvest_salt(pf, SideDir(cfg, Registry(cfg, pf))) == "deadbeef"
    (cfg / "harvest.salt").write_text("someone else's file\n")  # anything else is foreign and never replaced
    with pytest.raises(PreflightError):
        harvest_salt(pf, SideDir(cfg, Registry(cfg, pf)))


def test_control_characters_do_not_crash_the_writer(tmp_path):
    """N27: characters XML cannot carry are dropped wherever text enters the document."""
    from texdocxdrafter.harvest import normalise
    from texdocxdrafter.writer.xmlutil import run, xml_safe

    assert xml_safe("a\x00b\x01c\x0bd\te\nf") == "abcd\te\nf"
    assert normalise("x\x02y") == "xy"
    assert run("a\x00b").find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t").text == "ab"
    src = _write(tmp_path, "\\documentclass{article}\n\\title{T\x01itle}\n\\begin{document}\n\\maketitle\nbad \x00 char\n"
                 "\\footnote{n\x1fote}\n\\end{document}\n")
    res = build(src, no_harvest=True, log=lambda *_: None)
    assert res.check.ok, res.check.errors
    assert "bad  char" in _docx_xml(tmp_path / "main.docx") or "bad char" in _docx_xml(tmp_path / "main.docx")


# ---- ninth review round --------------------------------------------------------------------------

LUA_RECOVER = r"""\documentclass{article}
\begin{document}
\directlua{
  local dbg = require("debug")
  texio.write_nl("GETUPVALUE:" .. tostring(dbg.getupvalue) .. ":" .. tostring(package.loaded["debug"] == debug))
  local recovered
  if dbg.getupvalue then
    for i = 1, 20 do
      local name, val = dbg.getupvalue(io.open, i)
      if name == "io_open" then recovered = val end
    end
  end
  if recovered then
    local f = recovered("%(victim)s", "w") if f then f:write("pwned") f:close() end
  end
  local e1, m1 = os.execute("touch %(marker)s")
  local e2, m2 = io.popen("touch %(marker)s")
  local e3, m3 = package.loadlib("%(lib)s", "luaopen_os")
  local ok, mod = pcall(require, "texdocxdrafter_planted")
  texio.write_nl("EXEC:" .. tostring(e1) .. " POPEN:" .. tostring(e2) .. " LOADLIB:" .. tostring(e3) .. " REQUIRE:" .. tostring(ok))
}
x
\end{document}
"""


@pytest.mark.skipif(not HAVE_LUALATEX, reason="lualatex not installed")
def test_lua_guard_originals_cannot_be_recovered_and_native_code_cannot_load(tmp_path, monkeypatch):
    """N30, N31: require("debug") gives the restricted table; os.execute, io.popen, loadlib and C modules are off."""
    import glob

    monkeypatch.setitem(sandbox._disabled, "bwrap", True)
    victim = tmp_path / "victim.txt"
    victim.write_text("untouched")
    work = tmp_path / "work"
    work.mkdir()
    libs = glob.glob("/usr/lib/*/liblua5.*.so*") + glob.glob("/usr/lib/*/libm.so*")
    lib = libs[0] if libs else "/nonexistent.so"
    (work / "texdocxdrafter_planted.so").write_bytes(b"\x7fELF not really")
    tex = work / "evil.tex"
    tex.write_text(LUA_RECOVER % {"victim": str(victim), "marker": str(tmp_path / "ran.txt"), "lib": lib})
    r = sandbox.compile_tex("lualatex", work, "evil", tex, [work], timeout=120, check_policy=False)
    assert r.returncode != -1, r.stderr
    log = (work / "evil.log").read_text(errors="replace").replace("\n", "")
    assert "GETUPVALUE:nil:true" in log, log[-600:]
    assert "EXEC:nil POPEN:nil LOADLIB:nil REQUIRE:false" in log, log[-600:]
    assert victim.read_text() == "untouched" and not (tmp_path / "ran.txt").exists()


def test_malformed_numbering_ids_in_template_do_not_crash(tmp_path):
    """N32: a template with a non-numeric w:abstractNumId or w:numId still works as a reference document."""
    import docx

    tpl = tmp_path / "template.docx"
    docx.Document().save(str(tpl))
    with zipfile.ZipFile(tpl) as z:
        numbering = z.read("word/numbering.xml")
    first_child = numbering.index(b">", numbering.index(b"<w:numbering")) + 1
    bad = (b'<w:abstractNum w:abstractNumId="abc"><w:lvl w:ilvl="0"><w:numFmt w:val="bullet"/></w:lvl></w:abstractNum>'
           b'<w:num w:numId="x1"><w:abstractNumId w:val="abc"/></w:num>')
    _rezip(tpl, {"word/numbering.xml": [(numbering, numbering[:first_child] + bad + numbering[first_child:])]}, {})
    src = _write(tmp_path, DOC.replace("Hello world.", "\\begin{itemize}\\item one\\end{itemize}"))
    res = build(src, reference_doc=tpl, no_harvest=True, log=lambda *_: None)
    assert res.check.ok, res.check.errors


def _template_with_header_links(path: Path) -> dict:
    """A template whose header carries hyperlinks of every kind: w:hyperlink elements bound to relationships,
    HYPERLINK fields (simple and complex) and a bookmark link; and a body hyperlink relationship."""
    import docx
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import parse_xml

    W = ('xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
         'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"')
    d = docx.Document()
    d.add_paragraph("body")
    d.part.relate_to("file:///C:/secret.docx", RT.HYPERLINK, is_external=True)
    hdr = d.sections[0].header
    hdr.paragraphs[0].text = "Running head "
    p = hdr.paragraphs[0]._p
    rids = {}
    for key, url in (("web", "https://journal.example/guide"), ("mail", "mailto:editor@example.org"),
                     ("unc", "\\\\attacker\\share\\x"), ("js", "javascript:alert(1)"),
                     ("file", "file:///C:/Users/me/notes.docx"), ("rel", "../../hidden.docx")):
        rids[key] = hdr.part.relate_to(url, RT.HYPERLINK, is_external=True)
        p.append(parse_xml(f'<w:hyperlink {W} r:id="{rids[key]}"><w:r><w:t>LINK-{key}</w:t></w:r></w:hyperlink>'))
    for xml in (
        f'<w:hyperlink {W} w:anchor="top"><w:r><w:t>LINK-anchor</w:t></w:r></w:hyperlink>',
        f'<w:fldSimple {W} w:instr=" HYPERLINK &quot;https://journal.example/&quot; "><w:r><w:t>FLD-web</w:t></w:r></w:fldSimple>',
        f'<w:fldSimple {W} w:instr=" HYPERLINK &quot;\\\\\\\\attacker\\\\share&quot; "><w:r><w:t>FLD-unc</w:t></w:r></w:fldSimple>',
        f'<w:fldSimple {W} w:instr=" HYPERLINK &quot;javascript:alert(1)&quot; "><w:r><w:t>FLD-js</w:t></w:r></w:fldSimple>',
        f'<w:fldSimple {W} w:instr=" HYPERLINK \\l &quot;top&quot; "><w:r><w:t>FLD-bookmark</w:t></w:r></w:fldSimple>',
        f'<w:r {W}><w:fldChar w:fldCharType="begin"/></w:r>',
        f'<w:r {W}><w:instrText xml:space="preserve"> HYPERLINK "file:///C:/evil.hta" </w:instrText></w:r>',
        f'<w:r {W}><w:fldChar w:fldCharType="separate"/></w:r>',
        f'<w:r {W}><w:t>FLD-complex-file</w:t></w:r>',
        f'<w:r {W}><w:fldChar w:fldCharType="end"/></w:r>',
        f'<w:r {W}><w:fldChar w:fldCharType="begin"/></w:r>',
        f'<w:r {W}><w:instrText xml:space="preserve"> HYPERLINK "http://journal.example/x" </w:instrText></w:r>',
        f'<w:r {W}><w:fldChar w:fldCharType="separate"/></w:r>',
        f'<w:r {W}><w:t>FLD-complex-web</w:t></w:r>',
        f'<w:r {W}><w:fldChar w:fldCharType="end"/></w:r>',
    ):
        p.append(parse_xml(xml))
    d.save(str(path))
    return rids


def test_template_header_hyperlinks_keep_only_web_and_mail_targets(tmp_path):
    """N33: hyperlinks kept from a template's header obey the writer's scheme rule (http, https, mailto)."""
    tpl = tmp_path / "template.docx"
    _template_with_header_links(tpl)
    src = _write(tmp_path)
    res = build(src, reference_doc=tpl, no_harvest=True, log=lambda *_: None)
    assert res.check.ok, res.check.errors
    z = zipfile.ZipFile(tmp_path / "main.docx")
    names = z.namelist()
    headers = "".join(z.read(n).decode("utf-8") for n in names if n.startswith("word/header") and n.endswith(".xml"))
    rels = "".join(z.read(n).decode("utf-8") for n in names if n.startswith("word/_rels/header"))
    # every link text survives, as text; only web and mail targets are still links
    for key in ("web", "mail", "unc", "js", "file", "rel", "anchor"):
        assert f"LINK-{key}" in headers, key
    assert headers.count("<w:hyperlink") == 3  # web, mail, anchor
    assert "https://journal.example/guide" in rels and "mailto:editor@example.org" in rels
    for bad in ("attacker", "javascript", "notes.docx", "hidden.docx", "evil.hta", "secret.docx"):
        assert bad not in rels and bad not in headers, bad
    assert "FLD-web" in headers and "FLD-bookmark" in headers and "FLD-complex-web" in headers
    for gone in ("FLD-unc", "FLD-js", "FLD-complex-file"):
        assert gone not in headers, gone
    assert headers.count('w:fldCharType="begin"') == headers.count('w:fldCharType="end"') == 1
    doc_rels = z.read("word/_rels/document.xml.rels").decode("utf-8")
    assert "secret.docx" not in doc_rels  # the template body's own external links do not linger in the package


def test_unclosed_flag_conditionals_resolve_in_linear_time():
    """N34: a regex per \\newif flag with a lazy body rescanned to the end at every \\if it could not close."""
    from texdocxdrafter.pipeline import _resolve_flags

    assert _resolve_flags(r"a \ifdraft B \else C \fi d", {"draft": True}).split() == ["a", "B", "d"]
    assert _resolve_flags(r"a \ifdraft B \else C \fi d", {"draft": False}).split() == ["a", "C", "d"]
    assert _resolve_flags(r"x \iffalse gone \fi y", {}).split() == ["x", "y"]
    assert _resolve_flags(r"\ifdraft \ifnotes X \else W \fi Y \else Z \fi", {"draft": True, "notes": False}).split() == ["W", "Y"]
    n = 100_000
    for body, flags in ((r"\ifdraft word " * n, {"draft": True}), (r"\ifdraft word " * n, {"draft": False}),
                        (r"\iffalse word " * n, {}), (r"\ifa \ifb w " * n, {"a": True, "b": True})):
        _bounded(lambda: _resolve_flags(body, flags), 10)


def test_paragraph_of_many_line_breaks_strips_in_linear_time():
    """N35: leading line breaks were popped from the front of the list one by one."""
    from texdocxdrafter.parser.core import strip_inlines
    from texdocxdrafter.parser.model import LineBreak, Text

    items = [LineBreak() for _ in range(600_000)] + [Text(text=" x ")] + [LineBreak() for _ in range(600_000)]
    out = _bounded(lambda: strip_inlines(items), 10)
    assert len(out) == 1 and out[0].text == "x"
