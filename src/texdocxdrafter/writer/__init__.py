"""Stage 8: semantic tree → docx. One pass; python-docx for the package, lxml for fields, bookmarks, OMML."""
from __future__ import annotations

import hashlib
import re
from collections import Counter
from pathlib import Path

from docx.opc.constants import CONTENT_TYPE as CT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.shared import Inches, Mm
from docx.text.paragraph import Paragraph
from lxml import etree

from ..audit import Audit
from ..parser.model import (
    Bibliography, Cell, Cite, DisplayMath, Document, Figure, Footnote, Heading, InlineMath, Label, LineBreak,
    Link, ListBlock, PageBreak, Para, Quote, RawBlock, RawInline, Ref, Revision, RevisionBlock, Styled, Table, Text,
    Verbatim,
)
from ..refdoc import build_reference
from . import images as gfx
from . import math as omml
from .textmath import text_math
from .xmlutil import NSMAP, br, el, field, fld_char, instr, para, q, rpr, run, style_id

COLORS = {"red": "FF0000", "blue": "0000FF", "green": "008000", "black": "000000", "gray": "808080",
          "grey": "808080", "orange": "FFA500", "purple": "800080", "cyan": "00FFFF", "magenta": "FF00FF",
          "yellow": "FFFF00", "white": "FFFFFF"}
CAPTION_WORD = {"table": "Table", "figure": "Figure"}
REF_WORDS = {  # (autoref/cref word singular, plural) by kind
    "table": ("table", "tables"), "figure": ("figure", "figures"), "subfigure": ("figure", "figures"),
    "equation": ("equation", "equations"), "item": ("item", "items"), "para": ("paragraph", "paragraphs"),
}
SECTION_WORDS = {1: "section", 2: "subsection", 3: "subsubsection", 4: "paragraph", 5: "subparagraph"}
FOOT_CITES = {"footcite", "footcitetext", "footfullcite", "footcites"}


def bm_name(prefix: str, key: str) -> str:
    safe = re.sub(r"[^A-Za-z0-9_]", "_", key)
    name = f"{prefix}{safe}"
    if len(name) > 40:
        name = name[:31] + hashlib.md5(key.encode()).hexdigest()[:9]
    return name


class DocxWriter:
    def __init__(self, resolved, pf, sm: dict, audit: Audit, template=None, link_citations: bool = False):
        self.res = resolved
        self.doc: Document = resolved.doc
        self.labels = resolved.labels
        self.pf = pf
        self.sm = sm
        self.audit = audit
        self.link_citations = link_citations
        self.d, self.nm, self.refinfo = build_reference(sm, template, audit)
        self.body = self.d.element.body
        self.sect = self.body.find(q("w:sectPr"))
        self.sid = {s.name: s.style_id for s in self.d.styles}
        self.container = None
        self.override: str | None = None
        self._bm = 100
        self.seq = Counter()
        self.footnotes: list = []
        self.used = Counter()
        self.last_was_table = False
        self.heading_count = Counter()
        self.rev_stack: list[str] = []
        cols = self.sm["options"].get("columns", "auto")
        self.columns = (2 if pf.two_column else 1) if cols == "auto" else int(cols)
        self.cur_cols = 1
        self.breaks = 0
        self._rev_id = 100000

    # ---- infrastructure ----------------------------------------------------------------------
    def add(self, element) -> None:
        if self.container is None:
            self.sect.addprevious(element)
        else:
            self.container.append(element)
        self.last_was_table = element.tag == q("w:tbl")

    def sty(self, name: str) -> str:
        self.used[name] += 1
        return self.sid.get(name) or style_id(name)

    def ps(self, key: str) -> str:
        return self.sty(self.sm["paragraph"][key])

    def cs(self, key: str) -> str:
        return self.sid.get(self.sm["character"][key]) or style_id(self.sm["character"][key])

    def next_bm(self) -> int:
        self._bm += 1
        return self._bm

    def bookmark(self, name: str, inner: list) -> list:
        i = self.next_bm()
        return [el("w:bookmarkStart", {"w:id": i, "w:name": name}), *inner, el("w:bookmarkEnd", {"w:id": i})]

    def label_bm(self, key: str) -> str:
        return bm_name("_wf_", key)

    # ---- runs and inlines --------------------------------------------------------------------
    def rformat(self, flags: frozenset, link: bool = False) -> tuple[str | None, dict]:
        e, s, c = "emph" in flags, "strong" in flags, "code" in flags
        props: dict = {}
        rstyle = None
        if link:
            rstyle = self.cs("link")
            if e:
                props["i"] = True
            if s:
                props["b"] = True
        elif c:
            rstyle = self.cs("code")
            if e:
                props["i"] = True
            if s:
                props["b"] = True
        elif e and s:
            rstyle = self.cs("strong_emph")
        elif e:
            rstyle = self.cs("emph")
        elif s:
            rstyle = self.cs("strong")
        if "smallcaps" in flags:
            props["smallCaps"] = True
        if "underline" in flags:
            props["u"] = True
        if "strike" in flags:
            props["strike"] = True
        if "superscript" in flags:
            props["vertAlign"] = "superscript"
        if "subscript" in flags:
            props["vertAlign"] = "subscript"
        for f in flags:
            if isinstance(f, tuple) and f[0] == "color":
                c = f[1]
                props["color"] = COLORS.get(c.lower(), c.upper() if re.fullmatch(r"[0-9a-fA-F]{6}", c) else "000000")
        if rstyle:
            self.used[f"char:{rstyle}"] += 1
        return rstyle, props

    def text_run(self, text: str, flags: frozenset, link: bool = False):
        rs, props = self.rformat(flags, link)
        if props and not rs:
            self.audit.count("direct formatting runs")
        return run(text, rs, props)

    def inl(self, items: list, flags: frozenset = frozenset(), link: bool = False) -> list:
        out: list = []
        items = _swap_space(items)
        for x in items:
            if isinstance(x, Text):
                out.append(self.text_run(x.text, flags, link))
            elif isinstance(x, Styled):
                if x.kind == "color":
                    nf = flags | {("color", x.arg or "black")}
                elif x.kind == "emph" and "emph" in flags:
                    nf = flags - {"emph"}
                else:
                    nf = flags | {x.kind}
                out.extend(self.inl(x.children, frozenset(nf), link))
            elif isinstance(x, Link):
                out.append(self.hyperlink(x, flags))
            elif isinstance(x, Footnote):
                out.append(self.footnote_ref(x))
            elif isinstance(x, Cite):
                out.extend(self.cite_elements(x))
            elif isinstance(x, Ref):
                out.extend(self.ref_elements(x, flags))
            elif isinstance(x, InlineMath):
                out.extend(self.inline_math(x, flags))
            elif isinstance(x, LineBreak):
                out.append(br())
            elif isinstance(x, Label):
                out.extend(self.bookmark(self.label_bm(x.key), []))
            elif isinstance(x, Revision):
                self.rev_stack.append(x.kind)
                try:
                    els = self.inl(x.children, flags, link)
                finally:
                    self.rev_stack.pop()
                out.extend(self.mark_revision(els, x.kind))
            elif isinstance(x, RawInline):
                out.append(run(x.latex, self.cs("todo")))
            else:
                out.append(run(str(x), self.cs("todo")))
        return out

    def hyperlink(self, link: Link, flags: frozenset):
        rid = self.d.part.relate_to(link.url, RT.HYPERLINK, is_external=True)
        h = el("w:hyperlink", {"r:id": rid})
        for r in self.inl(link.children, flags, link=True):
            h.append(r)
        return h

    # ---- citations ---------------------------------------------------------------------------
    def cite_elements(self, c: Cite) -> list:
        if c.cmd == "nocite":
            return []
        runs = c.runs
        if runs is None:
            keys = ", ".join(re.findall(r"\{([^{}]*)\}", c.raw)[-1:]) or "?"
            return [run(f"[{keys}]", self.cs("todo"))]
        els = []
        for r in runs:
            fmt = set(r.fmt)
            flags = set()
            if "i" in fmt:
                flags.add("emph")
            if "b" in fmt:
                flags.add("strong")
            if "sup" in fmt:
                flags.add("superscript")
            if "sub" in fmt:
                flags.add("subscript")
            if r.url:
                els.append(self.hyperlink(Link(url=r.url, children=[Text(text=r.text)]), frozenset(flags)))
                continue
            rs, props = self.rformat(frozenset(flags))
            if rs is None:
                rs = self.cs("citation")
            els.append(run(r.text, rs, props))
        if c.cmd in FOOT_CITES:
            return [self.footnote_from_elements(els)]
        if self.link_citations:
            keys = re.findall(r"\{([^{}]*)\}", c.raw)
            first = keys[-1].split(",")[0].strip() if keys else ""
            if first:
                h = el("w:hyperlink", {"w:anchor": bm_name("_wfbib_", first), "w:history": "1"})
                for e in els:
                    h.append(e)
                return [h]
        return els

    # ---- cross-references --------------------------------------------------------------------
    def ref_elements(self, r: Ref, flags: frozenset) -> list:
        out: list = []
        keys = r.keys or []
        infos = [self.labels.get(k) for k in keys]
        word_mode = r.cmd in ("cref", "Cref", "autoref", "vref")
        if r.cmd == "eqref":
            out.append(run("(", None))
        if word_mode and infos and infos[0] is not None:
            out.append(self.text_run(self.ref_word(r, infos), flags))
            out.append(self.text_run(" ", flags))
        for i, (k, info) in enumerate(zip(keys, infos)):
            if i > 0:
                sep = " and " if (i == len(keys) - 1 and len(keys) == 2) else (", and " if i == len(keys) - 1 else ", ")
                out.append(self.text_run(sep if word_mode or r.cmd != "eqref" else ", ", flags))
            if info is None:
                out.append(run("??", self.cs("todo")))
                continue
            out.extend(self.one_ref(r.cmd, info))
        if r.cmd == "eqref":
            out.append(run(")", None))
        return out

    def ref_word(self, r: Ref, infos) -> str:
        first = infos[0]
        plural = len(infos) > 1
        if first.kind == "section":
            w = SECTION_WORDS.get(first.level, "section")
            word = w + ("s" if plural else "")
        else:
            sg, pl = REF_WORDS.get(first.kind, ("item", "items"))
            word = pl if plural else sg
        if r.cmd == "Cref" or (r.cmd == "autoref" and first.kind not in ("equation", "section")) or \
                (r.cmd == "autoref" and first.kind == "section" and first.level == 1):
            word = word[:1].upper() + word[1:]
        return word

    def one_ref(self, cmd: str, info) -> list:
        bm = self.label_bm(info.key)
        cached = info.number or "?"
        if cmd == "pageref":
            return field(f"PAGEREF {bm} \\h", info.page or "?")
        if cmd == "nameref":
            return field(f"REF {bm} \\h", cached)
        if info.kind == "para":
            # LaTeX prints the number of the enclosing section for a label inside running text
            head = self.labels.get(info.heading) if info.heading else None
            if head is None or (head.number and info.number and head.number != info.number):
                return [run(cached)]
            return field(f"REF {self.label_bm(head.key)} \\n \\h", cached)
        sw = {"section": "\\n ", "item": "\\n "}.get(info.kind, "")
        return field(f"REF {bm} {sw}\\h", cached)

    # ---- math --------------------------------------------------------------------------------
    def inline_math(self, m: InlineMath, flags) -> list:
        segs = text_math(m.latex) if self.sm["options"].get("inline_math_as_text", True) else None
        if segs is not None:
            self.audit.count("inline math as text")
            out = []
            for text, italic, vert in segs:
                props = {"vertAlign": vert} if vert else None
                if italic:
                    out.append(run(text, self.cs("math_var"), props))
                elif vert:
                    out.append(run(text, None, props))
                else:
                    out.append(self.text_run(text, flags))
            return out
        try:
            node = omml.inline_math(m.latex)
            self.audit.count("inline math (OMML)")
            return [node]
        except Exception as exc:
            self.audit.add("math", f"inline math kept as LaTeX text (${m.latex}$): {str(exc)[:60]}", m.loc)
            return [run("$" + m.latex + "$", self.cs("math"))]

    # ---- footnotes ---------------------------------------------------------------------------
    def new_footnote(self):
        fid = len(self.footnotes) + 1
        fn = el("w:footnote", {"w:id": fid})
        self.footnotes.append(fn)
        ref = el("w:r", None, rpr(self.cs("footnote_ref")))
        ref.append(el("w:footnoteReference", {"w:id": fid}))
        return fn, ref

    def footnote_ref(self, f: Footnote):
        fn, ref = self.new_footnote()
        saved = (self.container, self.override)
        self.container, self.override = fn, "footnote"
        try:
            for b in f.blocks:
                self.write_block(b)
        finally:
            self.container, self.override = saved
        self._prefix_footnote(fn)
        if self.rev_stack:
            self.mark_inside(fn, self.rev_stack[-1], marks=True)
        return ref

    def footnote_from_elements(self, els: list):
        fn, ref = self.new_footnote()
        fn.append(para(self.ps("footnote"), els))
        self._prefix_footnote(fn)
        return ref

    def _prefix_footnote(self, fn) -> None:
        p = fn.find(q("w:p"))
        if p is None:
            p = para(self.ps("footnote"))
            fn.append(p)
        mark = el("w:r", None, rpr(self.cs("footnote_ref")))
        mark.append(el("w:footnoteRef"))
        ppr = p.find(q("w:pPr"))
        idx = (list(p).index(ppr) + 1) if ppr is not None else 0
        p.insert(idx, mark)
        p.insert(idx + 1, run(" "))

    def finish_footnotes(self) -> None:
        if not self.footnotes:
            return
        root = etree.Element(q("w:footnotes"), nsmap=NSMAP)
        for fid, typ in ((-1, "separator"), (0, "continuationSeparator")):
            f = el("w:footnote", {"w:type": typ, "w:id": fid})
            r = el("w:r")
            r.append(el("w:separator") if typ == "separator" else el("w:continuationSeparator"))
            f.append(para(None, [r]))
            root.append(f)
        for fn in self.footnotes:
            root.append(fn)
        blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
        part = Part(PackURI("/word/footnotes.xml"), CT.WML_FOOTNOTES, blob, self.d.part.package)
        self.d.part.relate_to(part, RT.FOOTNOTES)
        settings = self.d.settings.element
        fp = el("w:footnotePr", None, el("w:footnote", {"w:id": -1}), el("w:footnote", {"w:id": 0}))
        _insert_settings(settings, fp, after=("w:characterSpacingControl",))

    # ---- blocks ------------------------------------------------------------------------------
    def write_blocks(self, blocks: list) -> None:
        for b in blocks:
            self.write_block(b)

    def write_block(self, b) -> None:
        if isinstance(b, Heading):
            self.write_heading(b)
        elif isinstance(b, Para):
            self.write_para(b)
        elif isinstance(b, ListBlock):
            self.write_list(b)
        elif isinstance(b, Quote):
            saved = self.override
            self.override = "quote"
            self.write_blocks(b.blocks)
            self.override = saved
        elif isinstance(b, Verbatim):
            for line in (b.text.split("\n") or [""]):
                r = el("w:r")
                t = el("w:t", None, text=line)
                t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                r.append(t)
                self.add(para(self.ps("code"), [r]))
        elif isinstance(b, Table):
            self.write_table(b)
        elif isinstance(b, Figure):
            self.write_figure(b)
        elif isinstance(b, DisplayMath):
            self.write_display(b)
        elif isinstance(b, Bibliography):
            self.write_bibliography(b)
        elif isinstance(b, PageBreak):
            self.add(para(self.ps("page_break"), [br("page")]))
        elif isinstance(b, RevisionBlock):
            new = self.capture(lambda: self.write_blocks(b.blocks), b.kind)
            for e in new:
                self.mark_inside(e, b.kind, marks=True)
        elif isinstance(b, RawBlock):
            self.add(para(self.ps("body"), [run(b.latex, self.cs("todo"))]))
        else:
            self.audit.add("other", f"unhandled block {type(b).__name__}")

    def para_style_key(self, p: Para) -> str:
        if p.style:
            return p.style
        if self.override:
            return self.override
        return "first_paragraph" if p.first else "body"

    def write_para(self, p: Para, style_key: str | None = None, ppr_extra=()) -> None:
        key = style_key or self.para_style_key(p)
        kids: list = []
        for lab in p.labels:
            kids.extend(self.bookmark(self.label_bm(lab), []))
        kids.extend(self.inl(p.children))
        self.add(para(self.ps(key), kids, ppr_extra))

    def write_heading(self, h: Heading, style_key: str | None = None) -> None:
        name = f"{self.sm['paragraph']['heading']} {h.level}" if style_key is None else self.sm["paragraph"][style_key]
        kids = self.inl(h.children)
        if h.label:
            kids = self.bookmark(self.label_bm(h.label), kids)
        extra = []
        if h.starred or style_key:
            extra.append(el("w:numPr", None, el("w:ilvl", {"w:val": 0}), el("w:numId", {"w:val": 0})))
        else:
            self.heading_count[h.level] += 1
        self.add(para(self.sty(name), kids, extra))

    def write_list(self, lb: ListBlock) -> None:
        lvl = min(lb.level, 2)
        suffix = ["", "_2", "_3"][lvl]
        if lb.kind == "itemize":
            key, num = "bullet_list" + suffix, self.nm.bullet_num()
        elif lb.kind == "enumerate":
            key, num = "number_list" + suffix, self.nm.new_decimal_list(lb.label_spec, lvl)
        else:
            key, num = "description_list", None
        for item in lb.items:
            first = True
            blocks = item.blocks or [Para(children=[])]
            if item.term is not None and not isinstance(blocks[0], Para):
                blocks = [Para(children=[])] + blocks  # the label line of an item that opens with a nested list
            for b in blocks:
                if isinstance(b, Para):
                    kids: list = []
                    seen_labs: set = set()
                    for lab in ([item.label] if item.label and first else []) + list(b.labels):
                        if lab and lab not in seen_labs:
                            seen_labs.add(lab)
                            kids.extend(self.bookmark(self.label_bm(lab), []))
                    custom = lb.kind != "description" and item.term is not None
                    if first and lb.kind == "description" and item.term:
                        kids.extend(self.inl(item.term, frozenset({"strong"})))
                        kids.append(run(" "))
                    elif first and custom:
                        # \\item[label] in itemize/enumerate: the label replaces the bullet or number
                        kids.extend(self.inl(item.term))
                        kids.append(el("w:r", None, el("w:tab")))
                    kids.extend(self.inl(b.children))
                    extra = []
                    if num is not None:
                        extra.append(el("w:numPr", None, el("w:ilvl", {"w:val": lvl}),
                                        el("w:numId", {"w:val": 0 if custom or not first else num})))
                    if custom:
                        extra.append(el("w:ind", {"w:left": 1080 + 720 * lvl, "w:hanging": 720}))
                    self.add(para(self.ps(key), kids, extra))
                else:
                    self.write_block(b)
                first = False

    # ---- captions, tables, figures -----------------------------------------------------------
    def caption(self, kind: str, inlines: list | None, label: str | None, loc) -> None:
        self.seq[kind] += 1
        n = self.seq[kind]
        info = self.labels.get(label) if label else None
        cached = str(n)
        if info and info.number and info.number != cached:
            self.audit.add("numbering", f"{kind} '{label}': LaTeX number {info.number} but Word sequence "
                           f"position {n}", loc)
        seq = field(f"SEQ {CAPTION_WORD[kind]} \\* ARABIC", cached)
        if label:
            seq = self.bookmark(self.label_bm(label), seq)
        kids = [run(CAPTION_WORD[kind] + " ")] + seq + [run(": ")] + self.inl(inlines or [])
        self.add(para(self.ps("caption_" + kind), kids))

    def write_notes(self, notes: list, key: str) -> None:
        saved = self.override
        self.override = key
        for n in notes:
            self.write_para(n, style_key=key)
        self.override = saved

    def write_table(self, t: Table) -> None:
        pos = self.sm["options"]["caption_position"]["table"]
        if self.last_was_table:
            self.add(para(self.ps("body")))
        if t.caption is not None or t.label:
            if pos == "above":
                self.caption("table", t.caption, t.label, t.loc)
        self.build_table(t)
        if t.caption is not None or t.label:
            if pos == "below":
                self.caption("table", t.caption, t.label, t.loc)
        if t.notes:
            self.write_notes(t.notes, "table_note")

    def column_widths(self, t: Table, total: int, pad: int = 220) -> list[int]:
        """Twip widths: natural width from the longest cell text, p{} columns from their ratio, X columns flexible.

        Every column first gets the width of its longest word (when the table can hold them all); the rest of
        the width is shared in proportion to what each column wants beyond that."""
        ncols = len(t.colspec)
        longest_text = [0] * ncols
        longest_word = [0] * ncols
        for r in t.rows:
            c = 0
            for cell in r.cells:
                if cell.colspan == 1 and c < ncols and not cell.vmerge:
                    text = "".join(_plain(b) for b in cell.blocks)
                    longest_text[c] = max(longest_text[c], min(len(text), 80))
                    longest_word[c] = max([longest_word[c]] + [len(w) for w in text.split()])
                c += cell.colspan
        per_char = 100
        widths: list[float] = []
        flex: list[int] = []
        for i in range(ncols):
            wt = t.colweights[i] if i < len(t.colweights) else 0
            floor = longest_word[i] * per_char + pad
            if wt < 0:
                flex.append(i)
                widths.append(0)
            elif wt > 0:
                widths.append(max(wt * 567 + pad, floor))  # centimetres to twips
            else:
                widths.append(max(longest_text[i] * per_char + pad, floor, 600))
        if flex:
            rest = total - sum(widths)
            rest = max(rest, 0.3 * total)
            for i in flex:
                widths[i] = max(rest / len(flex), longest_word[i] * per_char + pad)
        floors = [longest_word[i] * per_char + pad for i in range(ncols)]
        if sum(floors) >= total or sum(widths) <= total:
            s = sum(widths) or 1
            return [int(total * w / s) for w in widths]
        extra = [max(w - f, 0) for w, f in zip(widths, floors)]
        room = total - sum(floors)
        se = sum(extra) or 1
        return [int(f + room * e / se) for f, e in zip(floors, extra)]

    def build_table(self, t: Table) -> None:
        nrows, ncols = len(t.rows), len(t.colspec)
        if nrows == 0 or ncols == 0:
            return
        tbl = self.d.add_table(rows=nrows, cols=ncols)
        tbl.style = self.sm["options"]["table_style"]
        tbl_el = tbl._tbl
        if self.container is not None:  # nested (cell, footnote): move out of the body
            tbl_el.getparent().remove(tbl_el)
            self.container.append(tbl_el)
        self.last_was_table = True
        total = self.text_width_twips()
        narrow = self.cur_cols > 1  # a table inside one column: tight cell margins, as LaTeX sets it
        widths = self.column_widths(t, total, pad=110 if narrow else 220)
        pr = tbl_el.tblPr
        for old in pr.findall(q("w:tblW")):
            pr.remove(old)
        _tblpr_put(pr, el("w:tblW", {"w:w": total, "w:type": "dxa"}))
        look = pr.find(q("w:tblLook"))
        if look is not None:
            pr.remove(look)
        pr.append(el("w:tblLook", {"w:val": "0420" if t.header_rows else "0400", "w:firstRow": int(bool(t.header_rows)),
                                   "w:lastRow": 0, "w:firstColumn": 0, "w:lastColumn": 0, "w:noHBand": 1,
                                   "w:noVBand": 1}))
        if narrow:
            _tblpr_put(pr, el("w:tblCellMar", None, el("w:left", {"w:w": 50, "w:type": "dxa"}),
                              el("w:right", {"w:w": 50, "w:type": "dxa"})))
        grid = tbl_el.find(q("w:tblGrid"))
        for g, w in zip(grid.findall(q("w:gridCol")), widths):
            g.set(q("w:w"), str(w))
        # merges first, then content
        for ri, row in enumerate(t.rows):
            col = 0
            for cell in row.cells:
                if cell.vmerge:
                    col += cell.colspan
                    continue
                if cell.colspan > 1 or cell.rowspan > 1:
                    a = tbl.cell(ri, col)
                    b = tbl.cell(min(ri + cell.rowspan - 1, nrows - 1), min(col + cell.colspan - 1, ncols - 1))
                    if a._tc is not b._tc:
                        a.merge(b)
                col += cell.colspan
        for ri, row in enumerate(t.rows):
            col = 0
            tr = tbl_el.tr_lst[ri]
            if ri < t.header_rows:
                trpr = tr.get_or_add_trPr()
                trpr.append(el("w:cantSplit"))
                trpr.append(el("w:tblHeader"))
            for cell in row.cells:
                tc = tbl.cell(ri, col)._tc
                w = sum(widths[col:col + cell.colspan])
                _set_tcw(tc, w)
                if not cell.vmerge:
                    self.fill_cell(tc, cell)
                col += cell.colspan
        self.add_table_end_paragraph_fix(tbl_el)

    def add_table_end_paragraph_fix(self, tbl_el) -> None:
        for tc in tbl_el.iter(q("w:tc")):
            kids = list(tc)
            if not kids or kids[-1].tag != q("w:p"):
                tc.append(para(self.ps("table_text")))

    def fill_cell(self, tc, cell: Cell) -> None:
        for p in tc.findall(q("w:p")):
            tc.remove(p)
        base = "table_header" if cell.header else "table_text"
        key = base if cell.align == "l" else base + {"c": "_center", "r": "_right"}.get(cell.align, "")
        saved = (self.container, self.override, self.last_was_table)
        self.container, self.override = tc, key
        try:
            for b in cell.blocks:
                if isinstance(b, Para):
                    self.write_para(b, style_key=key)
                else:
                    self.write_block(b)
        finally:
            self.container, self.override, self.last_was_table = saved
        if tc.find(q("w:p")) is None:
            tc.append(para(self.ps(key)))

    COL_SPACE = 360  # twips between columns (0.25 in)

    def text_width_twips(self) -> int:
        """Width available to the block being written: the text block, or one column of it."""
        a4 = "a4paper" in self.pf.class_options
        full = 9026 if a4 else 9360
        return full if self.cur_cols == 1 else (full - self.COL_SPACE * (self.cur_cols - 1)) // self.cur_cols

    def write_figure(self, f: Figure) -> None:
        pos = self.sm["options"]["caption_position"]["figure"]
        has_cap = f.caption is not None or bool(f.label)
        if has_cap and pos == "above":
            self.caption("figure", f.caption, f.label, f.loc)
        multi = len(f.images) > 1
        for img in f.images:
            self.write_image(img)
            if multi and img.caption:
                self.add(para(self.ps("figure_note"), [run("(" + _panel(f.images.index(img)) + ") ")] +
                              self.inl(img.caption)))
        if has_cap and pos == "below":
            self.caption("figure", f.caption, f.label, f.loc)
        if f.notes:
            self.write_notes(f.notes, "figure_note")

    def write_image(self, img) -> None:
        if img.tikz is not None:
            if img.path is None:
                self.audit.add("raw", "tikzpicture not rendered: kept as visible text", img.loc)
                self.add(para(self.ps("body"), [run(img.tikz, self.cs("todo"))]))
                return
            self.add_picture(Path(img.path), img.width_in)
            return
        path = gfx.find_graphic(img.name, self.pf.srcdir, self.pf.graphicspath,
                                [self.pf.srcdir] + list(self.pf.allow_dirs))
        if path is None:
            self.audit.add("figure", f"graphic '{img.name}' not found", img.loc)
            self.add(para(self.ps("figure"), [run(f"[missing figure: {img.name}]", self.cs("todo"))]))
            return
        ras = gfx.to_raster(path, self.pf.workdir / "images")
        if ras is None:
            self.audit.add("figure", f"graphic '{path.name}' could not be converted", img.loc)
            self.add(para(self.ps("figure"), [run(f"[figure not converted: {path.name}]", self.cs("todo"))]))
            return
        from docx.image.image import Image as DImage

        try:
            im = DImage.from_file(str(ras))
            natural = im.px_width / (im.horz_dpi or 96)
        except Exception:
            natural = None
        avail = self.text_width_twips() / 1440
        w_in = gfx.width_inches(img.options, natural, avail)
        if w_in is None:
            w_in = min(natural or avail, avail)
        self.add_picture(ras, w_in)

    def add_picture(self, ras: Path, w_in: float | None) -> None:
        avail = self.text_width_twips() / 1440
        w_in = min(w_in or avail, avail)
        from docx.oxml import parse_xml

        p = parse_xml(etree.tostring(para(self.ps("figure"))))
        self.add(p)
        Paragraph(p, self.d._body).add_run().add_picture(str(ras), width=Inches(w_in))
        self.audit.count("figures embedded")

    # ---- columns -----------------------------------------------------------------------------
    FRONT_STYLES = ("title", "author", "date", "keywords")

    def block_columns(self, b) -> int:
        """Columns of the section that holds a top-level block: the title block and wide floats span the page."""
        if self.columns == 1:
            return 1
        if isinstance(b, Para) and b.style in self.FRONT_STYLES:
            return 1
        if isinstance(b, (Table, Figure)) and b.wide:
            return 1
        return self.columns

    def set_columns(self, sect, n: int, continuous: bool) -> None:
        cols = sect.find(q("w:cols"))
        if cols is None:
            cols = el("w:cols")
            anchor = sect.find(q("w:pgMar"))
            if anchor is not None:
                anchor.addnext(cols)
            else:
                sect.append(cols)
        cols.set(q("w:num"), str(n))
        cols.set(q("w:space"), str(self.COL_SPACE))
        typ = sect.find(q("w:type"))
        if continuous:
            if typ is None:
                typ = el("w:type")
                pgsz = sect.find(q("w:pgSz"))
                if pgsz is not None:
                    pgsz.addprevious(typ)
                else:
                    sect.insert(0, typ)
            typ.set(q("w:val"), "continuous")

    COMPAT_HEAD = ["useSingleBorderforContiguousCells", "wpJustification", "noTabHangInd", "noLeading", "spaceForUL"]

    def no_column_balance(self) -> None:
        """Columns fill top to bottom before a continuous break, the way a table* waits for the next page.

        Balanced columns ending at a break that also carry footnotes make LibreOffice 26.2 fail to lay the
        document out (it reports the file as corrupt), so the multi-column output leaves them unbalanced."""
        settings = self.d.settings.element
        compat = settings.find(q("w:compat"))
        if compat is None:
            compat = el("w:compat")
            _insert_settings(settings, compat)
        if compat.find(q("w:noColumnBalance")) is not None:
            return
        pos = 0
        for i, c in enumerate(compat):
            if etree.QName(c).localname in self.COMPAT_HEAD:
                pos = i + 1
        compat.insert(pos, el("w:noColumnBalance"))

    def section_break(self) -> None:
        """End the current section (with its column count) in a minimal empty paragraph."""
        import copy

        sp = copy.deepcopy(self.sect)
        self.set_columns(sp, self.cur_cols, continuous=self.breaks > 0)
        ppr = el("w:pPr", None, el("w:spacing", {"w:before": 0, "w:after": 0, "w:line": 20, "w:lineRule": "exact"}))
        ppr.append(sp)
        self.add(el("w:p", None, ppr))
        self.breaks += 1

    # ---- tracked changes ---------------------------------------------------------------------
    def capture(self, fn, kind: str) -> list:
        """Run fn with a revision kind active; return the top-level elements it added."""
        parent = self.container if self.container is not None else self.body
        before = len(parent)
        self.rev_stack.append(kind)
        try:
            fn()
        finally:
            self.rev_stack.pop()
        new = list(parent)[before - (1 if self.container is None else 0):]
        return [e for e in new if e is not self.sect]

    def rev_el(self, kind: str):
        self._rev_id += 1
        return el("w:ins" if kind == "ins" else "w:del",
                  {"w:id": self._rev_id, "w:author": self.sm["options"].get("revision_author", "LaTeX")})

    def mark_revision(self, els: list, kind: str) -> list:
        """Inline revision: wrap runs (and math runs) in w:ins/w:del; paragraph marks are untouched."""
        out = []
        for e in els:
            if e.tag in (q("w:r"), q("m:r")) or (kind == "ins" and e.tag == q("w:del")):
                w = self.rev_el(kind)
                w.append(e)
                self._as_deleted(e, kind)
                out.append(w)
            else:
                self.mark_inside(e, kind, marks=False)
                out.append(e)
        return out

    def mark_inside(self, root, kind: str, marks: bool) -> None:
        """Wrap every run under root; with marks, also mark paragraph marks and table rows."""
        for r in list(root.iter(q("w:r"), q("m:r"))):
            parent = r.getparent()
            # already marked; deleted text inside an insertion nests as w:ins > w:del > w:r
            if parent.tag == q("w:del") and kind == "ins" and parent.getparent().tag != q("w:ins"):
                r = parent
                parent = r.getparent()
            elif parent.tag in (q("w:del"), q("w:ins")) and (parent.tag == q("w:del") or kind == "ins"):
                continue
            w = self.rev_el(kind)
            parent.insert(parent.index(r), w)
            w.append(r)
            self._as_deleted(r, kind)
        if not marks:
            return
        for p in root.iter(q("w:p")):
            ppr = p.find(q("w:pPr"))
            if ppr is None:
                ppr = el("w:pPr")
                p.insert(0, ppr)
            rp = ppr.find(q("w:rPr"))
            if rp is None:
                rp = el("w:rPr")
                ppr.append(rp)
            if rp.find(q("w:ins")) is None and rp.find(q("w:del")) is None:
                rp.insert(0, self.rev_el(kind))
        for tr in root.iter(q("w:tr")):
            trpr = tr.find(q("w:trPr"))
            if trpr is None:
                trpr = el("w:trPr")
                pex = tr.find(q("w:tblPrEx"))
                tr.insert(1 if pex is not None else 0, trpr)
            trpr.append(self.rev_el(kind))

    @staticmethod
    def _as_deleted(r, kind: str) -> None:
        if kind != "del":
            return
        for t in r.iter(q("w:t")):
            t.tag = q("w:delText")
        for t in r.iter(q("w:instrText")):
            t.tag = q("w:delInstrText")

    # ---- comments ----------------------------------------------------------------------------
    def write_comments(self, anchors: list) -> None:
        """Turn LaTeX % comments into Word comments on the paragraph they annotate."""
        if not self.pf.comments or not anchors:
            return
        self.comment_styles()
        src = self.pf.source
        ranked = []
        last = -1
        for loc, p in anchors:
            r = src.rank(loc)
            last = r if r >= 0 else last
            ranked.append((last, p))
        groups: list[tuple[int, str, list[str], object]] = []
        prev = None
        for c in self.pf.comments:
            r = src.rank((c.file, c.line))
            if (prev is not None and c.attach == "next" and prev.attach == "next" and prev.file == c.file
                    and prev.line == c.line - 1):
                groups[-1][2].append(c.text)
            else:
                groups.append((r, c.attach, [c.text], c))
            prev = c
        for r, attach, texts, c in groups:
            if attach == "prev":
                cands = [p for k, p in ranked if k <= r]
                p = cands[-1] if cands else ranked[0][1]
            else:
                cands = [p for k, p in ranked if k > r]
                p = cands[0] if cands else ranked[-1][1]
            par = Paragraph(p, self.d._body)
            runs = par.runs or [par.add_run("")]
            self.d.add_comment([runs[0], runs[-1]], text="\n".join(texts), author="LaTeX", initials="TeX")
            self.audit.count("LaTeX comments written as Word comments")

    def comment_styles(self) -> None:
        """Word's built-in comment styles, added only when the template lacks them."""
        from docx.oxml import parse_xml

        styles = self.d.styles.element
        have = {s.get(q("w:styleId")) for s in styles.iter(q("w:style"))}
        ns = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
        if "CommentText" not in have:
            styles.append(parse_xml(
                f'<w:style {ns} w:type="paragraph" w:styleId="CommentText"><w:name w:val="annotation text"/>'
                '<w:basedOn w:val="Normal"/><w:uiPriority w:val="99"/><w:unhideWhenUsed/>'
                '<w:rPr><w:sz w:val="20"/><w:szCs w:val="20"/></w:rPr></w:style>'))
        if "CommentReference" not in have:
            styles.append(parse_xml(
                f'<w:style {ns} w:type="character" w:styleId="CommentReference"><w:name w:val="annotation reference"/>'
                '<w:uiPriority w:val="99"/><w:unhideWhenUsed/>'
                '<w:rPr><w:sz w:val="16"/><w:szCs w:val="16"/></w:rPr></w:style>'))

    # ---- equations ---------------------------------------------------------------------------
    def write_display(self, m: DisplayMath) -> None:
        lines = m.lines
        if not lines:
            return
        if not any(l.numbered for l in lines):
            if len(lines) > 1:
                try:
                    node = omml.eq_array([l.latex for l in lines])
                    self.add(para(self.ps("equation"), [omml.omath_para(node)]))
                    self.audit.count("display math (OMML)")
                    return
                except Exception as exc:
                    self.audit.add("math", f"{m.env}: {str(exc)[:60]}; lines converted separately", m.loc, "info")
            for l in lines:
                self.add(para(self.ps("equation"), self.display_runs(l, m)))
            return
        # numbered: borderless layout table, one row per line
        tbl = self.d.add_table(rows=len(lines), cols=3)
        tbl.style = self.sm["options"]["equation_table_style"]
        tbl_el = tbl._tbl
        if self.container is not None:
            tbl_el.getparent().remove(tbl_el)
            self.container.append(tbl_el)
        self.last_was_table = True
        total = self.text_width_twips()
        widths = [int(total * 0.08), int(total * 0.84), int(total * 0.08)]
        pr = tbl_el.tblPr
        for old in pr.findall(q("w:tblW")):
            pr.remove(old)
        _tblpr_put(pr, el("w:tblW", {"w:w": total, "w:type": "dxa"}))
        _tblpr_put(pr, el("w:tblLayout", {"w:type": "fixed"}))
        for g, w in zip(tbl_el.find(q("w:tblGrid")).findall(q("w:gridCol")), widths):
            g.set(q("w:w"), str(w))
        for ri, l in enumerate(lines):
            cells = [tbl.cell(ri, c)._tc for c in range(3)]
            for tc, w in zip(cells, widths):
                _set_tcw(tc, w)
                for p in tc.findall(q("w:p")):
                    tc.remove(p)
            cells[0].append(para(self.ps("equation_number")))
            cells[1].append(para(self.ps("equation"), self.display_runs(l, m)))
            if l.numbered:
                self.seq["Equation"] += 1
                n = self.seq["Equation"]
                info = self.labels.get(l.label) if l.label else None
                if info and info.number and info.number != str(n):
                    self.audit.add("numbering", f"equation '{l.label}': LaTeX number {info.number} but Word "
                                   f"sequence position {n}", l.loc)
                seq = field("SEQ Equation \\* ARABIC", str(n))
                if l.label:
                    seq = self.bookmark(self.label_bm(l.label), seq)
                cells[2].append(para(self.ps("equation_number"), [run("(")] + seq + [run(")")]))
            else:
                cells[2].append(para(self.ps("equation_number")))
        self.audit.count("numbered equations")

    def display_runs(self, line, m: DisplayMath) -> list:
        try:
            node = omml.convert(line.latex.replace("&", " "))
            self.audit.count("display math (OMML)")
            return [omml.omath_para(node)]
        except Exception as exc:
            self.audit.add("math", f"display math kept as LaTeX text: {str(exc)[:60]}", line.loc)
            return [run(line.latex, self.cs("math"))]

    # ---- bibliography ------------------------------------------------------------------------
    def write_bibliography(self, b: Bibliography) -> None:
        title = b.title
        if title is None:
            t = b.options.get("title")
            title = [Text(text=re.sub(r"[{}]", "", t))] if isinstance(t, str) and t else [Text(text="References")]
        h = Heading(level=1, children=title, starred=True, loc=b.loc)
        self.write_heading(h, style_key="bibliography_heading")
        for e in b.entries:
            els = []
            for r in e.runs:
                flags = set()
                if "i" in r.fmt:
                    flags.add("emph")
                if "b" in r.fmt:
                    flags.add("strong")
                if "sup" in r.fmt:
                    flags.add("superscript")
                if r.url:
                    els.append(self.hyperlink(Link(url=r.url, children=[Text(text=r.text)]), frozenset(flags)))
                else:
                    els.append(self.text_run(r.text, frozenset(flags)))
            els = self.bookmark(bm_name("_wfbib_", e.key), els)
            self.add(para(self.ps("bibliography"), els))
        if not b.entries:
            self.audit.add("harvest", "bibliography is empty", b.loc)

    # ---- document ----------------------------------------------------------------------------
    def page_setup(self) -> None:
        sec = self.d.sections[0]
        if "a4paper" in self.pf.class_options:
            sec.page_width, sec.page_height = Mm(210), Mm(297)
        else:
            sec.page_width, sec.page_height = Inches(8.5), Inches(11)
        sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Inches(1)

    def write(self, out, comments: bool = True):
        self.page_setup()
        anchors = []  # (block loc, first paragraph written for the block)
        for k, b in enumerate(self.doc.blocks):
            want = self.block_columns(b)
            if k == 0:
                self.cur_cols = want
            elif want != self.cur_cols:
                self.section_break()
                self.cur_cols = want
            before = self.sect.getprevious()
            self.write_block(b)
            new = before.getnext() if before is not None else self.body[0]
            if new is not None and new is not self.sect:
                p = new if new.tag == q("w:p") else next(new.iter(q("w:p")), None)
                if p is not None:
                    anchors.append((b.loc, p))
        self.set_columns(self.sect, self.cur_cols, continuous=self.breaks > 0)
        if self.columns > 1:
            self.audit.info["columns"] = f"{self.columns} ({self.breaks} section breaks)"
            if self.breaks:
                self.no_column_balance()
        if comments:
            self.write_comments(anchors)
        self.finish_footnotes()
        self.nm.apply(self.d.part.numbering_part.element)
        settings = self.d.settings.element
        _insert_settings(settings, el("w:updateFields", {"w:val": "true"}), after=("w:zoom", "w:removePersonalInformation"))
        cp = self.d.core_properties
        if self.doc.title:
            cp.title = "".join(_inline_text(self.doc.title))
        if self.doc.authors:
            cp.author = "; ".join("".join(_inline_text(a)) for a in self.doc.authors)
        cp.last_modified_by = "texdocxdrafter"  # lets a later run tell its own untouched output from an edited draft
        # a file object (the pipeline writes the bytes itself, atomically) or a path
        self.d.save(out if hasattr(out, "write") else str(out))
        self.audit.info["styles used"] = len(self.used)
        return out


def _panel(i: int) -> str:
    return chr(ord("a") + i)


def _plain(b) -> str:
    if isinstance(b, Para):
        return "".join(_inline_text(b.children))
    return ""


TBLPR_ORDER = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize", "tblW",
               "jc", "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook",
               "tblCaption", "tblDescription"]


def _tblpr_put(pr, child) -> None:
    """Insert a tblPr child at its schema position; Word rejects out-of-order table properties."""
    rank = TBLPR_ORDER.index(etree.QName(child).localname)
    for c in pr:
        name = etree.QName(c).localname
        if name in TBLPR_ORDER and TBLPR_ORDER.index(name) > rank:
            c.addprevious(child)
            return
    pr.append(child)


def _swap_space(items: list) -> list:
    """Drop a lone space between a deletion and an insertion (\\del{five} \\add{four}), so that accepting or
    rejecting every change leaves single spacing."""
    out = list(items)
    for i in range(len(out) - 2, 0, -1):
        a, sp, b = out[i - 1], out[i], out[i + 1]
        if (isinstance(sp, Text) and sp.text.isspace() and isinstance(a, Revision) and isinstance(b, Revision)
                and a.kind != b.kind):
            del out[i]
    return out


def _inline_text(items) -> list[str]:
    out = []
    for x in items:
        if isinstance(x, Text):
            out.append(x.text)
        elif isinstance(x, (Styled, Link)) or (isinstance(x, Revision) and x.kind == "ins"):
            out.extend(_inline_text(x.children))
        elif isinstance(x, InlineMath):
            out.append(x.latex)
        elif isinstance(x, LineBreak):
            out.append(" ")
    return out


def _set_tcw(tc, w: int) -> None:
    pr = tc.get_or_add_tcPr()
    for old in pr.findall(q("w:tcW")):
        pr.remove(old)
    pr.insert(0, el("w:tcW", {"w:w": w, "w:type": "dxa"}))


SETTINGS_ORDER = ["w:writeProtection", "w:view", "w:zoom", "w:removePersonalInformation", "w:removeDateAndTime",
                  "w:doNotDisplayPageBoundaries", "w:displayBackgroundShape", "w:printPostScriptOverText",
                  "w:printFractionalCharacterWidth", "w:printFormsData", "w:embedTrueTypeFonts",
                  "w:embedSystemFonts", "w:saveSubsetFonts", "w:saveFormsData", "w:mirrorMargins",
                  "w:alignBordersAndEdges", "w:bordersDoNotSurroundHeader", "w:bordersDoNotSurroundFooter",
                  "w:gutterAtTop", "w:hideSpellingErrors", "w:hideGrammaticalErrors", "w:activeWritingStyle",
                  "w:proofState", "w:formsDesign", "w:attachedTemplate", "w:linkStyles", "w:stylePaneFormatFilter",
                  "w:stylePaneSortMethod", "w:documentType", "w:mailMerge", "w:revisionView", "w:trackRevisions",
                  "w:doNotTrackMoves", "w:doNotTrackFormatting", "w:documentProtection", "w:autoFormatOverride",
                  "w:styleLockTheme", "w:styleLockQFSet", "w:defaultTabStop", "w:autoHyphenation",
                  "w:consecutiveHyphenLimit", "w:hyphenationZone", "w:doNotHyphenateCaps", "w:showEnvelope",
                  "w:summaryLength", "w:clickAndTypeStyle", "w:defaultTableStyle", "w:evenAndOddHeaders",
                  "w:bookFoldRevPrinting", "w:bookFoldPrinting", "w:bookFoldPrintingSheets",
                  "w:drawingGridHorizontalSpacing", "w:drawingGridVerticalSpacing",
                  "w:displayHorizontalDrawingGridEvery", "w:displayVerticalDrawingGridEvery",
                  "w:doNotUseMarginsForDrawingGridOrigin", "w:drawingGridHorizontalOrigin",
                  "w:drawingGridVerticalOrigin", "w:doNotShadeFormData", "w:noPunctuationKerning",
                  "w:characterSpacingControl", "w:printTwoOnOne", "w:strictFirstAndLastChars",
                  "w:noLineBreaksAfter", "w:noLineBreaksBefore", "w:savePreviewPicture",
                  "w:doNotValidateAgainstSchema", "w:saveInvalidXml", "w:ignoreMixedContent",
                  "w:alwaysShowPlaceholderText", "w:doNotDemarcateInvalidXml", "w:saveXmlDataOnly",
                  "w:useXSLTWhenSaving", "w:saveThroughXslt", "w:showXMLTags", "w:alwaysMergeEmptyNamespace",
                  "w:updateFields", "w:hdrShapeDefaults", "w:footnotePr", "w:endnotePr", "w:compat", "w:docVars",
                  "w:rsids", "m:mathPr", "w:attachedSchema", "w:themeFontLang", "w:clrSchemeMapping",
                  "w:doNotIncludeSubdocsInStats", "w:doNotAutoCompressPictures", "w:forceUpgrade", "w:captions",
                  "w:readModeInkLockDown", "w:smartTagType", "w:schemaLibrary", "w:shapeDefaults",
                  "w:doNotEmbedSmartTags", "w:decimalSymbol", "w:listSeparator"]


def _insert_settings(settings, new, after=()) -> None:
    """Insert ``new`` into settings.xml at the position the schema requires."""
    order = {q(t): i for i, t in enumerate(SETTINGS_ORDER)}
    for old in settings.findall(new.tag):
        settings.remove(old)
    rank = order.get(new.tag, len(order))
    pos = 0
    for i, child in enumerate(settings):
        if order.get(child.tag, len(order)) <= rank:
            pos = i + 1
    settings.insert(pos, new)
